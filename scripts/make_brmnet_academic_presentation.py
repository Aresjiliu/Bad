from __future__ import annotations

import argparse
import sys
from pathlib import Path


ROOT = Path(__file__).resolve().parents[1]
PPT_DEPS = ROOT / "output" / "ppt_deps"
if PPT_DEPS.exists():
    sys.path.insert(0, str(PPT_DEPS))

from pptx import Presentation  # type: ignore  # noqa: E402
from pptx.dml.color import RGBColor  # type: ignore  # noqa: E402
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION  # type: ignore  # noqa: E402
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE  # type: ignore  # noqa: E402
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN  # type: ignore  # noqa: E402
from pptx.util import Inches, Pt  # type: ignore  # noqa: E402
from pptx.chart.data import ChartData  # type: ignore  # noqa: E402


PAPER = Path("D:/Academic/paper_submission/brmnet_pricai2026")
FIG = PAPER / "figures" / "generated"
OUT = ROOT / "docs" / "slides" / "BRM-Net_Academic_Presentation.pptx"

NAVY = RGBColor(18, 35, 62)
NAVY2 = RGBColor(28, 57, 92)
CYAN = RGBColor(29, 165, 180)
CYAN_DARK = RGBColor(13, 118, 137)
ORANGE = RGBColor(231, 126, 52)
GRAY = RGBColor(102, 112, 128)
LIGHT = RGBColor(238, 244, 248)
OFFWHITE = RGBColor(249, 251, 252)
WHITE = RGBColor(255, 255, 255)
BLACK = RGBColor(20, 26, 33)
GREEN = RGBColor(60, 156, 119)
RED = RGBColor(188, 73, 73)

SLIDES: list[str] = []
TODOS = [
    "Confusion matrix or class-wise map for Houston2013 is not available as a final paper figure.",
    "Final latency table is not used because the current submission tables focus on MAC/parameter ratios.",
    "Dense continuous corruption curves beyond the current noise/downsample/occlusion protocol remain future work.",
]


def build_parser() -> argparse.ArgumentParser:
    p = argparse.ArgumentParser()
    p.add_argument("--output", default=str(OUT))
    return p


def rgb(hex_text: str) -> RGBColor:
    text = hex_text.strip("#")
    return RGBColor(int(text[:2], 16), int(text[2:4], 16), int(text[4:], 16))


def textbox(slide, text: str, x: float, y: float, w: float, h: float, size: int = 16,
            color=BLACK, bold: bool = False, align=PP_ALIGN.LEFT, font="Aptos",
            valign=MSO_ANCHOR.TOP):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = shape.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = Inches(0.03)
    tf.margin_right = Inches(0.03)
    tf.margin_top = Inches(0.02)
    tf.margin_bottom = Inches(0.02)
    tf.vertical_anchor = valign
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    return shape


def multiline(slide, lines: list[tuple[str, str]], x: float, y: float, w: float, h: float,
              size: int = 15, color=BLACK):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = shape.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = Inches(0.06)
    tf.margin_right = Inches(0.04)
    tf.margin_top = Inches(0.02)
    tf.margin_bottom = Inches(0.02)
    for i, (head, body) in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.level = 0
        p.space_after = Pt(6)
        r = p.add_run()
        r.text = head
        r.font.name = "Aptos"
        r.font.size = Pt(size)
        r.font.bold = True
        r.font.color.rgb = color
        r2 = p.add_run()
        r2.text = body
        r2.font.name = "Aptos"
        r2.font.size = Pt(size)
        r2.font.color.rgb = color
    return shape


def bullet_list(slide, items: list[str], x: float, y: float, w: float, h: float,
                size: int = 16, color=BLACK):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = shape.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = Inches(0.06)
    tf.margin_right = Inches(0.04)
    tf.margin_top = Inches(0.02)
    tf.margin_bottom = Inches(0.02)
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = item
        p.level = 0
        p.space_after = Pt(7)
        p.font.name = "Aptos"
        p.font.size = Pt(size)
        p.font.color.rgb = color
    return shape


def add_slide(prs: Presentation, title: str, section: str, dark: bool = False):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = NAVY if dark else OFFWHITE
    if not dark:
        bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(0.18))
        bar.fill.solid()
        bar.fill.fore_color.rgb = NAVY
        bar.line.fill.background()
        textbox(slide, section.upper(), 0.63, 0.31, 2.2, 0.24, size=8, color=CYAN_DARK, bold=True)
        textbox(slide, title, 0.63, 0.53, 11.2, 0.52, size=28, color=NAVY, bold=True)
    SLIDES.append(title)
    return slide


def footer(slide, idx: int):
    textbox(slide, "BRM-Net | academic presentation", 0.63, 7.12, 4.2, 0.18, size=8, color=GRAY)
    textbox(slide, f"{idx:02d}", 12.15, 7.06, 0.48, 0.2, size=10, color=GRAY, align=PP_ALIGN.RIGHT)


def card(slide, x: float, y: float, w: float, h: float, title: str, body: str,
         accent=CYAN, fill=WHITE):
    s = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    s.line.color.rgb = RGBColor(220, 228, 235)
    stripe = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(0.08), Inches(h))
    stripe.fill.solid()
    stripe.fill.fore_color.rgb = accent
    stripe.line.fill.background()
    textbox(slide, title, x + 0.2, y + 0.18, w - 0.35, 0.28, size=14, color=NAVY, bold=True)
    textbox(slide, body, x + 0.2, y + 0.58, w - 0.35, h - 0.66, size=11, color=GRAY)
    return s


def dark_card(slide, x: float, y: float, w: float, h: float, title: str, body: str, accent=CYAN):
    s = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    s.fill.solid()
    s.fill.fore_color.rgb = RGBColor(28, 52, 82)
    s.line.color.rgb = RGBColor(138, 160, 184)
    stripe = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(0.08), Inches(h))
    stripe.fill.solid()
    stripe.fill.fore_color.rgb = accent
    stripe.line.fill.background()
    textbox(slide, title, x + 0.2, y + 0.18, w - 0.35, 0.28, size=14, color=WHITE, bold=True)
    textbox(slide, body, x + 0.2, y + 0.58, w - 0.35, h - 0.66, size=11, color=LIGHT)
    return s


def kpi(slide, x: float, y: float, w: float, value: str, label: str, color=CYAN):
    textbox(slide, value, x, y, w, 0.52, size=30, color=color, bold=True, align=PP_ALIGN.CENTER)
    textbox(slide, label, x, y + 0.55, w, 0.3, size=10, color=GRAY, align=PP_ALIGN.CENTER)


def add_picture(slide, path: Path, x: float, y: float, w: float, h: float):
    if not path.exists():
        card(slide, x, y, w, h, "Missing figure", "待补充最终实验图表", RED)
        return None
    pic = slide.shapes.add_picture(str(path), Inches(x), Inches(y), width=Inches(w))
    if pic.height > Inches(h):
        ratio = pic.width / pic.height
        pic.height = Inches(h)
        pic.width = int(pic.height * ratio)
    pic.left = Inches(x + (w - pic.width / 914400) / 2)
    return pic


def simple_table(slide, headers: list[str], rows: list[list[str]], x: float, y: float, w: float, h: float,
                 font_size: int = 9):
    table = slide.shapes.add_table(len(rows) + 1, len(headers), Inches(x), Inches(y), Inches(w), Inches(h)).table
    for c, head in enumerate(headers):
        cell = table.cell(0, c)
        cell.text = head
        cell.fill.solid()
        cell.fill.fore_color.rgb = NAVY
        p = cell.text_frame.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        p.runs[0].font.color.rgb = WHITE
        p.runs[0].font.bold = True
        p.runs[0].font.size = Pt(font_size)
    for r, row in enumerate(rows, 1):
        for c, val in enumerate(row):
            cell = table.cell(r, c)
            cell.text = val
            cell.fill.solid()
            cell.fill.fore_color.rgb = WHITE if r % 2 else LIGHT
            p = cell.text_frame.paragraphs[0]
            p.alignment = PP_ALIGN.CENTER
            p.runs[0].font.size = Pt(font_size)
            p.runs[0].font.color.rgb = BLACK
    return table


def connector(slide, x1: float, y1: float, x2: float, y2: float, color=CYAN):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = color
    c.line.width = Pt(2)
    return c


def chart_bar(slide, cats, series: list[tuple[str, list[float]]], x, y, w, h, title: str | None = None):
    data = ChartData()
    data.categories = cats
    for name, values in series:
        data.add_series(name, values)
    frame = slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(x), Inches(y), Inches(w), Inches(h), data)
    chart = frame.chart
    chart.has_legend = True
    chart.legend.position = XL_LEGEND_POSITION.BOTTOM
    chart.legend.include_in_layout = False
    chart.value_axis.tick_labels.font.size = Pt(8)
    chart.category_axis.tick_labels.font.size = Pt(8)
    if title:
        chart.has_title = True
        chart.chart_title.text_frame.text = title
        chart.chart_title.text_frame.paragraphs[0].runs[0].font.size = Pt(11)
    return chart


def slide_01(prs):
    s = add_slide(prs, "Title", "opening", dark=True)
    textbox(s, "BRM-Net", 0.72, 0.72, 4.2, 0.45, size=21, color=CYAN, bold=True)
    textbox(s, "Fusion-Compatible Budgeted Compact Export\nwith Degradation-Supervised Reliability", 0.72, 1.45, 11.6, 1.18, size=33, color=WHITE, bold=True)
    textbox(s, "for Multimodal Remote Sensing Classification", 0.76, 2.85, 10.4, 0.35, size=18, color=LIGHT)
    dark_card(s, 0.78, 4.1, 3.4, 1.2, "Core question", "Can a multimodal RS model be physically compact and robust to missing/degraded modalities?", CYAN)
    dark_card(s, 4.6, 4.1, 3.4, 1.2, "Main evidence", "Actual MAC budget agreement + robustness under controlled modality states.", ORANGE)
    dark_card(s, 8.42, 4.1, 3.4, 1.2, "Report scope", "Submission-oriented story; thesis extensions remain future work.", GREEN)
    textbox(s, "Graduation pre-defense / group meeting / submission report", 0.78, 6.62, 7.2, 0.28, size=12, color=LIGHT)


def slide_02(prs):
    s = add_slide(prs, "Research motivation: why lightweight multimodal RS matters", "motivation")
    kpi(s, 0.95, 1.55, 2.1, "2+", "sensor branches", CYAN)
    kpi(s, 3.55, 1.55, 2.1, "80%", "target MAC setting", ORANGE)
    kpi(s, 6.15, 1.55, 2.1, "3", "dataset roles", GREEN)
    card(s, 8.95, 1.18, 3.2, 2.35, "Deployment pressure", "Edge/onboard inference needs lower MACs, parameters, memory traffic, and stable behavior.", ORANGE)
    bullet_list(s, [
        "HSI provides rich spectral signatures but is high-dimensional.",
        "LiDAR or auxiliary sensors add geometry and spatial cues.",
        "Multi-branch fusion improves accuracy but raises cost and failure modes.",
        "A deployable method must report measured compact resources, not only training-time gates.",
    ], 0.95, 4.05, 10.8, 1.6, size=17)
    footer(s, 2)


def slide_03(prs):
    s = add_slide(prs, "Problem: resource budget + missing/degraded modality", "problem")
    card(s, 0.8, 1.35, 3.3, 2.0, "Budget constraint", "The final network should satisfy a target resource ratio rho measured by actual MACs/params.", CYAN)
    card(s, 4.35, 1.35, 3.3, 2.0, "Missing modality", "A modality can be physically unavailable and must be excluded from fusion competition.", RED)
    card(s, 7.9, 1.35, 3.3, 2.0, "Degraded modality", "A modality can remain available but become noisy, low-resolution, or occluded.", ORANGE)
    textbox(s, "Desired behavior", 0.9, 4.25, 3.2, 0.3, size=19, color=NAVY, bold=True)
    bullet_list(s, [
        "Keep complete-modality accuracy competitive.",
        "Export a physically smaller compact model.",
        "Handle full, main-only, auxiliary-only, and corrupted auxiliary inputs.",
        "Treat reliability as a degradation-supervised diagnostic signal.",
    ], 0.95, 4.72, 10.2, 1.45, size=16)
    footer(s, 3)


def slide_04(prs):
    s = add_slide(prs, "Key gap in existing work", "gap")
    headers = ["Method family", "Budget", "Compact export", "Missing", "Degraded", "Fusion interface"]
    rows = [
        ["Structured pruning", "Yes", "Yes", "No", "No", "Usually single-backbone"],
        ["Elastic networks", "Yes", "Partial", "No", "No", "Width-adaptive only"],
        ["Multimodal RS fusion", "No", "No", "Limited", "Limited", "Accuracy-first"],
        ["Missing-modality RS", "Limited", "No", "Yes", "Partial", "Often heavier"],
        ["BRM-Net", "Yes", "Yes", "Yes", "Yes", "Explicitly preserved"],
    ]
    simple_table(s, headers, rows, 0.65, 1.45, 12.0, 2.5, font_size=8)
    card(s, 0.9, 4.55, 3.5, 1.15, "Gap 1", "Compression rarely respects multimodal fusion interfaces.", CYAN)
    card(s, 4.9, 4.55, 3.5, 1.15, "Gap 2", "Missing-modality methods often rely on reconstruction or heavier modules.", ORANGE)
    card(s, 8.9, 4.55, 3.5, 1.15, "Gap 3", "Degraded-but-available modalities are under-tested.", RED)
    footer(s, 4)


def slide_05(prs):
    s = add_slide(prs, "BRM-Net overview", "method")
    xs = [0.7, 3.2, 5.7, 8.2, 10.7]
    titles = ["Inputs", "Gated encoders", "Reliability", "Masked fusion", "Compact export"]
    bodies = ["HSI/main\nAuxiliary", "Hard-concrete\nchannel gates", "Quality scores\nq_m", "Availability-aware\nsoftmax", "Physical channel\nremoval"]
    colors = [NAVY, CYAN, ORANGE, GREEN, NAVY2]
    for i, x in enumerate(xs):
        card(s, x, 2.0, 2.0, 1.65, titles[i], bodies[i], colors[i])
        if i < len(xs) - 1:
            connector(s, x + 2.0, 2.83, x + 2.45, 2.83, CYAN)
    textbox(s, "Training", 1.4, 4.35, 1.3, 0.28, size=15, color=CYAN_DARK, bold=True)
    connector(s, 2.45, 4.48, 5.15, 4.48, CYAN)
    textbox(s, "Compact Export", 5.25, 4.35, 2.0, 0.28, size=15, color=ORANGE, bold=True)
    connector(s, 7.0, 4.48, 9.25, 4.48, ORANGE)
    textbox(s, "Inference", 9.35, 4.35, 1.5, 0.28, size=15, color=GREEN, bold=True)
    textbox(s, "Main design principle: compress internal channels while preserving the terminal fusion feature interface.", 1.1, 5.7, 10.8, 0.4, size=18, color=NAVY, bold=True, align=PP_ALIGN.CENTER)
    footer(s, 5)


def slide_06(prs):
    s = add_slide(prs, "Problem formulation", "method")
    textbox(s, "Constrained objective", 0.9, 1.38, 3.3, 0.36, size=20, color=NAVY, bold=True)
    eq = "min E[L_cls(f_theta(g(x), a), y)]\nsubject to R(f_theta^c) <= rho R(f_theta)"
    card(s, 0.92, 1.9, 4.8, 1.7, "Deployment target", eq, CYAN)
    multiline(s, [
        ("x^(1), x^(2): ", "heterogeneous aligned observations such as HSI-LiDAR."),
        ("a_m in {0,1}: ", "availability mask; missing modalities are excluded from fusion."),
        ("g in G: ", "controlled degradation: noise, downsampling, occlusion."),
        ("f_theta^c: ", "physically compact model after structural projection."),
    ], 6.25, 1.42, 5.7, 2.35, size=14)
    card(s, 1.1, 4.55, 10.9, 1.25, "Narrow claim", "BRM-Net does not claim physical sensor calibration. It evaluates degradation-supervised reliability as a diagnostic fusion signal under controlled modality states.", ORANGE)
    footer(s, 6)


def slide_07(prs):
    s = add_slide(prs, "Budget-aware channel gating", "method")
    card(s, 0.85, 1.35, 3.35, 2.1, "Gate definition", "Each gated layer has z_l in [0,1]^C. Features are multiplied channel-wise before export.", CYAN)
    card(s, 4.75, 1.35, 3.35, 2.1, "Training estimate", "Hard-concrete probabilities estimate expected channel usage and resource cost.", ORANGE)
    card(s, 8.65, 1.35, 3.35, 2.1, "Export rule", "Threshold gates are binarized to meet target MACs while respecting minimum widths.", GREEN)
    chart_bar(s, ["65", "80", "90"], [("Target", [65, 80, 90]), ("Actual", [64.96, 80.01, 90.07])], 1.2, 4.1, 5.2, 2.2, "MAC budget agreement")
    bullet_list(s, [
        "Resource loss penalizes mismatch between expected and target MAC ratio.",
        "Gating is only an intermediate training mechanism.",
        "The final report evaluates exported compact models.",
    ], 7.0, 4.15, 4.8, 1.55, size=15)
    footer(s, 7)


def slide_08(prs):
    s = add_slide(prs, "Fusion-compatible compact export", "method")
    card(s, 0.85, 1.35, 3.3, 1.65, "Why compatibility matters", "Independent branch pruning can break the feature dimensions expected by the fusion head.", RED)
    card(s, 4.55, 1.35, 3.3, 1.65, "Structural projection", "Remove output channels, dependent input slices, BN parameters, and classifier weights.", CYAN)
    card(s, 8.25, 1.35, 3.3, 1.65, "Terminal constraint", "Preserve fusion-consumed terminal representations through tied or protected gates.", GREEN)
    simple_table(s, ["State", "Source OA", "Compact OA", "Delta", "MACs"], [
        ["Full", "86.15", "86.90", "+0.75", "80.01"],
        ["HSI only", "78.18", "80.04", "+1.86", "80.01"],
        ["Noise-high", "84.26", "84.55", "+0.29", "80.01"],
        ["Occlusion-50", "84.59", "85.57", "+0.98", "80.01"],
    ], 1.0, 4.0, 10.9, 1.65, font_size=9)
    textbox(s, "Source gated model and physically exported compact model remain close, supporting the compact-export claim.", 1.05, 6.1, 10.8, 0.3, size=14, color=NAVY, bold=True)
    footer(s, 8)


def slide_09(prs):
    s = add_slide(prs, "Availability-aware reliability fusion", "method")
    card(s, 0.85, 1.35, 3.45, 2.0, "Reliability scores", "Each modality produces q_m, a learned quality score used in fusion weighting.", ORANGE)
    card(s, 4.72, 1.35, 3.45, 2.0, "Availability mask", "Unavailable modalities receive a large negative logit before softmax.", CYAN)
    card(s, 8.59, 1.35, 3.45, 2.0, "Fusion weights", "alpha_m = softmax(log q_m + mask(a_m)) over usable modalities.", GREEN)
    simple_table(s, ["State", "BRM-Net", "w/o mask", "Delta"], [
        ["Full", "86.90", "87.02", "+0.12"],
        ["HSI only", "80.04", "76.72", "-3.32"],
        ["LiDAR only", "38.04", "37.59", "-0.45"],
        ["Occlusion-50", "85.57", "85.67", "+0.10"],
    ], 1.2, 4.15, 10.2, 1.55, font_size=10)
    textbox(s, "Interpretation: the mask mainly protects physically missing modalities, not every degraded-but-available case.", 1.15, 6.1, 10.6, 0.28, size=14, color=NAVY, bold=True)
    footer(s, 9)


def slide_10(prs):
    s = add_slide(prs, "Degradation-supervised quality learning", "method")
    card(s, 0.8, 1.25, 2.65, 1.5, "Clean / available", "target t_m = 1", GREEN)
    card(s, 3.75, 1.25, 2.65, 1.5, "Missing", "target t_m = 0", RED)
    card(s, 6.7, 1.25, 2.65, 1.5, "Noise / occlusion", "intermediate target", ORANGE)
    card(s, 9.65, 1.25, 2.65, 1.5, "Downsample x4", "lower target", ORANGE)
    textbox(s, "Training loss", 0.95, 3.75, 2.0, 0.3, size=19, color=NAVY, bold=True)
    card(s, 0.95, 4.2, 5.05, 1.05, "Objective", "L = L_cls + lambda_b L_bud + lambda_q L_qual", CYAN)
    bullet_list(s, [
        "Degraded modalities remain available during fusion.",
        "The reliability branch must learn to down-weight corrupted inputs.",
        "This is not a calibrated physical sensor-quality estimator.",
    ], 6.65, 3.85, 5.2, 1.45, size=15)
    footer(s, 10)


def slide_11(prs):
    s = add_slide(prs, "Training and export pipeline", "method")
    phases = [
        ("1", "Source training", "Modality dropout + degradation sampling + budget loss"),
        ("2", "Gate projection", "Select threshold eta to match target MAC ratio"),
        ("3", "Compact export", "Remove channels and dependent parameters"),
        ("4", "Fine-tune", "Recover accuracy in the compact architecture"),
        ("5", "Evaluate", "Full, missing, degraded, MACs, params, latency"),
    ]
    for i, (num, title, body) in enumerate(phases):
        x = 0.75 + i * 2.48
        card(s, x, 1.65, 2.08, 2.15, f"{num}. {title}", body, [CYAN, ORANGE, GREEN, NAVY2, CYAN_DARK][i])
        if i < len(phases) - 1:
            connector(s, x + 2.08, 2.72, x + 2.42, 2.72, CYAN)
    card(s, 1.15, 4.85, 10.9, 1.1, "Algorithmic emphasis", "The export stage is essential: the deployment model has physically fewer channels rather than merely soft gates in the source network.", ORANGE)
    footer(s, 11)


def slide_12(prs):
    s = add_slide(prs, "Experimental protocol and datasets", "experiments")
    simple_table(s, ["Dataset", "Role", "Train/Test", "Why included"], [
        ["Houston2013", "Main benchmark", "2,832 / 12,197", "Official split; full ablation protocol"],
        ["Trento", "External validation", "819 / 29,395", "Near-saturated HSI-LiDAR accuracy test"],
        ["MUUFL", "Stress test", "1,550 / 52,137", "Imbalanced scene; weak auxiliary modality"],
    ], 0.75, 1.4, 11.9, 2.0, font_size=9)
    card(s, 0.9, 4.2, 3.35, 1.35, "Metrics", "OA, AA, Kappa, MAC ratio, parameter ratio, latency where available.", CYAN)
    card(s, 4.95, 4.2, 3.35, 1.35, "States", "Full, main-only, auxiliary-only, noise, downsampling, occlusion.", ORANGE)
    card(s, 9.0, 4.2, 3.35, 1.35, "Reporting rule", "Use fixed seeds and avoid leaderboard-style cross-dataset overclaiming.", GREEN)
    footer(s, 12)


def slide_13(prs):
    s = add_slide(prs, "Budget agreement results", "results")
    kpi(s, 1.05, 1.4, 2.0, "64.96%", "actual MAC at 65 target", CYAN)
    kpi(s, 4.0, 1.4, 2.0, "80.01%", "actual MAC at 80 target", ORANGE)
    kpi(s, 6.95, 1.4, 2.0, "90.07%", "actual MAC at 90 target", GREEN)
    chart_bar(s, ["65", "80", "90"], [("Actual MAC", [64.96, 80.01, 90.07]), ("Compact OA", [85.53, 85.84, 85.74])], 1.15, 3.5, 5.7, 2.45, "Budget and OA")
    simple_table(s, ["Target", "Source OA", "Compact OA", "Params", "MACs"], [
        ["65", "85.98", "85.53", "64.79", "64.96"],
        ["80", "85.79", "85.84", "79.92", "80.01"],
        ["90", "85.33", "85.74", "89.78", "90.07"],
    ], 7.25, 3.55, 4.55, 1.75, font_size=9)
    textbox(s, "Main claim: target budgets can be materialized as measured compact networks.", 1.0, 6.28, 10.8, 0.3, size=15, color=NAVY, bold=True)
    footer(s, 13)


def slide_14(prs):
    s = add_slide(prs, "Accuracy-efficiency trade-off", "results")
    add_picture(s, FIG / "fig_pareto_accuracy_efficiency.png", 0.85, 1.25, 5.25, 4.75)
    add_picture(s, FIG / "fig_brmnet_results_overview.png", 6.35, 1.25, 5.75, 4.75)
    textbox(s, "Evidence combines Pareto-style accuracy-efficiency view and the paper overview figure.", 0.95, 6.18, 10.8, 0.32, size=14, color=NAVY, bold=True)
    footer(s, 14)


def slide_15(prs):
    s = add_slide(prs, "Missing-modality robustness", "results")
    simple_table(s, ["Training", "Full", "HSI only", "LiDAR only", "LiDAR noise"], [
        ["no dropout", "86.74", "63.34", "32.27", "86.13"],
        ["dropout 0.25", "87.01", "75.84", "42.95", "84.09"],
    ], 0.95, 1.35, 5.6, 1.15, font_size=10)
    chart_bar(s, ["Full", "HSI only", "LiDAR only", "Noise"], [
        ("no dropout", [86.74, 63.34, 32.27, 86.13]),
        ("dropout 0.25", [87.01, 75.84, 42.95, 84.09]),
    ], 0.95, 3.0, 5.9, 2.75, "Modality dropout effect")
    card(s, 7.35, 1.4, 4.2, 1.35, "Observation", "Training-time modality dropout greatly improves single-modality fallback.", GREEN)
    card(s, 7.35, 3.15, 4.2, 1.35, "Boundary", "Availability masking protects missing modalities; degradation needs quality supervision.", ORANGE)
    footer(s, 15)


def slide_16(prs):
    s = add_slide(prs, "Degraded-modality robustness", "results")
    simple_table(s, ["Method", "Full", "Noise-high", "Down-4", "Occ-50", "Adverse"], [
        ["Full baseline", "86.97", "72.08", "78.37", "74.21", "74.89"],
        ["Uniform fusion", "86.72", "77.01", "78.35", "73.87", "76.41"],
        ["Noise quality", "87.45", "85.15", "80.14", "73.32", "79.54"],
        ["Multi-deg p=.50", "85.99", "85.11", "84.38", "84.66", "84.72"],
        ["Multi-deg p=.25", "86.90", "84.55", "84.15", "85.57", "84.76"],
    ], 0.65, 1.28, 12.0, 2.15, font_size=8)
    chart_bar(s, ["Baseline", "Uniform", "Noise", "Multi .50", "Multi .25"], [
        ("Adverse Avg.", [74.89, 76.41, 79.54, 84.72, 84.76]),
        ("Full OA", [86.97, 86.72, 87.45, 85.99, 86.90]),
    ], 1.15, 4.0, 6.1, 2.35, "Clean vs adverse states")
    card(s, 7.85, 4.2, 3.9, 1.3, "Key result", "Adverse-state average improves from 74.89 to 84.76 under the 80% MAC setting.", ORANGE)
    footer(s, 16)


def slide_17(prs):
    s = add_slide(prs, "Multi-dataset evidence: Houston / Trento / MUUFL", "results")
    simple_table(s, ["Dataset", "Full OA", "Main-only", "Occlusion-50", "MACs"], [
        ["Houston2013", "86.90 ± 1.65", "80.04 ± 1.93", "85.57 ± 1.64", "80.01"],
        ["Trento", "98.94 ± 0.94", "92.89 ± 3.15", "98.16 ± 1.87", "79.90"],
        ["MUUFL", "86.87 ± 1.79", "83.03 ± 2.14", "85.99 ± 2.08", "79.97"],
    ], 0.75, 1.3, 11.75, 1.65, font_size=10)
    chart_bar(s, ["Houston", "Trento", "MUUFL"], [
        ("Full OA", [86.90, 98.94, 86.87]),
        ("Main-only", [80.04, 92.89, 83.03]),
        ("Occlusion-50", [85.57, 98.16, 85.99]),
    ], 1.0, 3.55, 6.0, 2.65, "Cross-dataset role evidence")
    card(s, 7.65, 3.7, 3.9, 1.65, "Interpretation", "Houston supports the main ablation story; Trento supports efficiency preservation; MUUFL exposes clean-robustness trade-off.", CYAN)
    footer(s, 17)


def slide_18(prs):
    s = add_slide(prs, "Ablation study", "results")
    simple_table(s, ["Ablation", "What is removed/changed", "Observed effect"], [
        ["Uniform fusion", "No reliability weighting", "Limited corruption robustness"],
        ["Noise quality", "Only noise-supervised quality", "Noise improves; occlusion remains weak"],
        ["Multi-deg p=.50", "Stronger degradation schedule", "Robust but slightly hurts full OA"],
        ["Multi-deg p=.25", "Light multi-degradation schedule", "Best adverse-state average"],
        ["w/o mask", "Disable availability mask", "HSI-only OA drops by 3.32 pp"],
    ], 0.75, 1.35, 11.85, 2.4, font_size=9)
    card(s, 0.9, 4.55, 5.0, 1.35, "Conclusion", "Robustness comes from degradation-supervised reliability, not from a reliability scalar alone.", ORANGE)
    card(s, 6.65, 4.55, 5.0, 1.35, "Submission boundary", "Routing, weighted CE, and profile label stabilization remain thesis extensions.", CYAN)
    footer(s, 18)


def slide_19(prs):
    s = add_slide(prs, "Visualization: retained width / reliability diagnostic / MUUFL error", "visualization")
    add_picture(s, FIG / "fig_gate_retention.png", 0.75, 1.25, 3.35, 4.55)
    add_picture(s, FIG / "fig_controlled_corruption_reliability_curves.png", 4.35, 1.25, 4.35, 4.55)
    add_picture(s, FIG / "fig_muufl_error_analysis.png", 8.95, 1.25, 3.45, 4.55)
    textbox(s, "Visuals should be read as mechanism diagnostics; reliability is degradation-aware, not physically calibrated.", 0.95, 6.22, 10.8, 0.3, size=14, color=NAVY, bold=True)
    footer(s, 19)


def slide_20(prs):
    s = add_slide(prs, "Discussion: clean-robustness trade-off and limitations", "discussion")
    card(s, 0.85, 1.35, 3.4, 1.65, "Supported", "Target MAC agreement and compact export are well supported on Houston2013.", GREEN)
    card(s, 4.8, 1.35, 3.4, 1.65, "Nuanced", "MUUFL improves robustness but sacrifices 1.64 pp clean full-modality OA.", ORANGE)
    card(s, 8.75, 1.35, 3.4, 1.65, "Limited", "Current datasets are still HSI-LiDAR style; HS-MS and HS-SAR remain future evidence.", RED)
    bullet_list(s, [
        "A single full-modality OA is not enough for deployment-oriented evaluation.",
        "Reliability weighting is useful only when trained with corrupted-but-available states.",
        "The paper should avoid claims of real onboard deployment or physical sensor calibration.",
        "Larger backbones and stronger lightweight baselines would strengthen submission impact.",
    ], 1.0, 4.0, 10.8, 1.65, size=16)
    footer(s, 20)


def slide_21(prs):
    s = add_slide(prs, "Contribution summary", "closing")
    kpi(s, 0.95, 1.45, 2.0, "1", "compact export", CYAN)
    kpi(s, 3.6, 1.45, 2.0, "2", "reliability fusion", ORANGE)
    kpi(s, 6.25, 1.45, 2.0, "3", "joint evaluation", GREEN)
    card(s, 0.95, 3.55, 3.4, 1.6, "Fusion-compatible budgeted compact export", "Learns gates and materializes compact models while preserving the fusion interface.", CYAN)
    card(s, 4.85, 3.55, 3.4, 1.6, "Degradation-supervised reliability fusion", "Uses quality targets for corrupted-but-available modalities without reconstruction modules.", ORANGE)
    card(s, 8.75, 3.55, 3.4, 1.6, "Accuracy-efficiency-robustness evaluation", "Reports actual MACs, missing-modality states, degraded states, and multi-dataset evidence.", GREEN)
    footer(s, 21)


def slide_22(prs):
    s = add_slide(prs, "Future work and Q&A", "closing", dark=True)
    textbox(s, "Future work", 0.75, 0.72, 3.5, 0.45, size=28, color=WHITE, bold=True)
    dark_card(s, 0.85, 1.65, 3.45, 1.45, "Extended thesis exploration", "Dynamic routing, weighted CE, profile label stabilization, class-wise diagnostics.", CYAN)
    dark_card(s, 4.85, 1.65, 3.45, 1.45, "Stronger evidence", "Larger-backbone transfer, stronger lightweight baselines, denser corruption curves.", ORANGE)
    dark_card(s, 8.85, 1.65, 3.45, 1.45, "Demo system", "Interactive dataset/state/budget selector for thesis defense demonstration.", GREEN)
    textbox(s, "Q&A", 0.85, 4.45, 3.0, 0.55, size=34, color=CYAN, bold=True)
    textbox(s, "Main message: BRM-Net targets deployable compact multimodal classification under budget, missing-modality, and degraded-modality constraints.", 0.9, 5.25, 10.8, 0.6, size=18, color=WHITE, bold=True)
    textbox(s, "TODO: final confusion matrix / final latency table / denser corruption grid remain outside this PPT's confirmed evidence.", 0.9, 6.45, 10.8, 0.25, size=11, color=LIGHT)


def build(output: Path) -> None:
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    for fn in [
        slide_01, slide_02, slide_03, slide_04, slide_05, slide_06, slide_07, slide_08,
        slide_09, slide_10, slide_11, slide_12, slide_13, slide_14, slide_15, slide_16,
        slide_17, slide_18, slide_19, slide_20, slide_21, slide_22,
    ]:
        fn(prs)
    output.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(output))


def main() -> int:
    args = build_parser().parse_args()
    output = Path(args.output)
    build(output)
    print(f"Wrote {output}")
    print(f"Slides: {len(SLIDES)}")
    for i, title in enumerate(SLIDES, 1):
        print(f"{i:02d}. {title}")
    print("TODO:")
    for item in TODOS:
        print(f"- {item}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
