from __future__ import annotations

import argparse
import sys
from pathlib import Path


ROOT = Path(__file__).resolve().parents[1]
PPT_DEPS = ROOT / "output" / "ppt_deps"
if PPT_DEPS.exists():
    sys.path.insert(0, str(PPT_DEPS))

from pptx import Presentation  # type: ignore  # noqa: E402
from pptx.dml.color import RGBColor  # type: ignore  # noqa: E402
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION  # type: ignore  # noqa: E402
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE  # type: ignore  # noqa: E402
from pptx.enum.text import PP_ALIGN  # type: ignore  # noqa: E402
from pptx.util import Inches, Pt  # type: ignore  # noqa: E402
from pptx.chart.data import ChartData  # type: ignore  # noqa: E402


PAPER = Path("D:/Academic/paper_submission/brmnet_pricai2026")
FIG = PAPER / "figures" / "generated"
OUT = ROOT / "docs" / "slides" / "BRM-Net_Academic_Presentation_ZH_v3.pptx"
FORMULA_DIR = ROOT / "docs" / "slides" / "formula_assets"

NAVY = RGBColor(18, 35, 62)
BLUE = RGBColor(29, 92, 145)
CYAN = RGBColor(28, 166, 183)
ORANGE = RGBColor(231, 126, 52)
GREEN = RGBColor(55, 151, 116)
RED = RGBColor(190, 72, 72)
GRAY = RGBColor(92, 103, 118)
LIGHT = RGBColor(235, 242, 247)
OFFWHITE = RGBColor(249, 251, 252)
WHITE = RGBColor(255, 255, 255)
BLACK = RGBColor(18, 25, 34)

SLIDES: list[str] = []


def parser() -> argparse.ArgumentParser:
    p = argparse.ArgumentParser()
    p.add_argument("--output", default=str(OUT))
    return p


def ensure_formulas() -> dict[str, Path]:
    import matplotlib.pyplot as plt

    FORMULA_DIR.mkdir(parents=True, exist_ok=True)
    formulas = {
        "objective": (
            r"$\min_{\theta}\ \mathbb{E}_{(x,y),a,g}"
            r"\left[\mathcal{L}_{\mathrm{cls}}\left(f_{\theta}(g(x),a),y\right)\right],"
            r"\quad R(f_{\theta}^{c})\leq \rho R(f_{\theta})$"
        ),
        "gates": (
            r"$\mathbf{z}^{(\ell)}\in[0,1]^{C_\ell},\qquad"
            r"\widetilde{\mathbf{F}}^{(\ell)}_{c,:,:}=z^{(\ell)}_c\mathbf{F}^{(\ell)}_{c,:,:}$"
        ),
        "fusion": (
            r"$\alpha_m=\frac{a_m\exp(q_m/\tau)}{\sum_{j=1}^{M}a_j\exp(q_j/\tau)+\epsilon},"
            r"\qquad \mathbf{F}_{\mathrm{fuse}}=\sum_{m=1}^{M}\alpha_m\widetilde{\mathbf{F}}_m$"
        ),
        "loss": (
            r"$\mathcal{L}=\mathcal{L}_{\mathrm{cls}}+\lambda_b\mathcal{L}_{\mathrm{bud}}"
            r"+\lambda_q\mathcal{L}_{\mathrm{qual}}$"
        ),
    }
    paths: dict[str, Path] = {}
    for key, formula in formulas.items():
        path = FORMULA_DIR / f"{key}.png"
        fig = plt.figure(figsize=(12, 1.25), dpi=300)
        fig.patch.set_alpha(0)
        plt.axis("off")
        plt.text(0.5, 0.5, formula, ha="center", va="center", fontsize=22, color="#12233e")
        plt.savefig(path, transparent=True, bbox_inches="tight", pad_inches=0.08)
        plt.close(fig)
        paths[key] = path
    return paths


def slide(prs: Presentation, title: str, section: str = "", dark: bool = False):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bg = s.background.fill
    bg.solid()
    bg.fore_color.rgb = NAVY if dark else OFFWHITE
    if not dark:
        top = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(0.16))
        top.fill.solid()
        top.fill.fore_color.rgb = NAVY
        top.line.fill.background()
        text(s, section, 0.62, 0.34, 2.8, 0.22, 10, CYAN, True)
        text(s, title, 0.62, 0.58, 11.8, 0.5, 31, NAVY, True)
    SLIDES.append(title)
    return s


def text(s, val: str, x: float, y: float, w: float, h: float, size: int = 20,
         color=BLACK, bold: bool = False, align=PP_ALIGN.LEFT):
    box = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = Inches(0.02)
    tf.margin_right = Inches(0.02)
    tf.margin_top = 0
    tf.margin_bottom = 0
    p = tf.paragraphs[0]
    p.alignment = align
    r = p.add_run()
    r.text = val
    r.font.name = "Microsoft YaHei"
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.color.rgb = color
    return box


def bullets(s, vals: list[str], x: float, y: float, w: float, h: float, size: int = 20):
    box = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = Inches(0.04)
    for i, val in enumerate(vals):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = val
        p.font.name = "Microsoft YaHei"
        p.font.size = Pt(size)
        p.font.color.rgb = BLACK
        p.space_after = Pt(10)
    return box


def box(s, x: float, y: float, w: float, h: float, title: str, body: str,
        accent=CYAN, fill=WHITE, body_size: int = 18):
    shp = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shp.fill.solid()
    shp.fill.fore_color.rgb = fill
    shp.line.color.rgb = RGBColor(210, 222, 232)
    stripe = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(0.09), Inches(h))
    stripe.fill.solid()
    stripe.fill.fore_color.rgb = accent
    stripe.line.fill.background()
    text(s, title, x + 0.22, y + 0.18, w - 0.36, 0.34, 20, NAVY, True)
    text(s, body, x + 0.22, y + 0.64, w - 0.36, h - 0.72, body_size, GRAY)


def dark_box(s, x, y, w, h, title, body, accent=CYAN):
    shp = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shp.fill.solid()
    shp.fill.fore_color.rgb = RGBColor(30, 55, 86)
    shp.line.color.rgb = RGBColor(130, 154, 180)
    stripe = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(0.09), Inches(h))
    stripe.fill.solid()
    stripe.fill.fore_color.rgb = accent
    stripe.line.fill.background()
    text(s, title, x + 0.22, y + 0.18, w - 0.36, 0.34, 20, WHITE, True)
    text(s, body, x + 0.22, y + 0.68, w - 0.36, h - 0.78, 18, LIGHT)


def module_card(s, x: float, y: float, w: float, h: float, label: str, accent=CYAN, size: int = 20):
    shp = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shp.fill.solid()
    shp.fill.fore_color.rgb = WHITE
    shp.line.color.rgb = RGBColor(210, 222, 232)
    stripe = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(0.08), Inches(h))
    stripe.fill.solid()
    stripe.fill.fore_color.rgb = accent
    stripe.line.fill.background()
    tf = shp.text_frame
    tf.clear()
    tf.margin_left = Inches(0.15)
    tf.margin_right = Inches(0.08)
    tf.margin_top = Inches(0.08)
    tf.margin_bottom = Inches(0.08)
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = label
    r.font.name = "Microsoft YaHei"
    r.font.size = Pt(size)
    r.font.bold = True
    r.font.color.rgb = NAVY
    return shp


def flow_card(s, x: float, y: float, w: float, h: float, title: str, lines: list[str], accent=CYAN):
    shp = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shp.fill.solid()
    shp.fill.fore_color.rgb = WHITE
    shp.line.color.rgb = RGBColor(210, 222, 232)
    stripe = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(0.09), Inches(h))
    stripe.fill.solid()
    stripe.fill.fore_color.rgb = accent
    stripe.line.fill.background()
    tf = shp.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = Inches(0.18)
    tf.margin_right = Inches(0.12)
    tf.margin_top = Inches(0.12)
    tf.margin_bottom = Inches(0.12)
    p = tf.paragraphs[0]
    p.space_after = Pt(8)
    r = p.add_run()
    r.text = title
    r.font.name = "Microsoft YaHei"
    r.font.size = Pt(21)
    r.font.bold = True
    r.font.color.rgb = NAVY
    for line_text in lines:
        p = tf.add_paragraph()
        p.space_after = Pt(2)
        r = p.add_run()
        r.text = line_text
        r.font.name = "Microsoft YaHei"
        r.font.size = Pt(18)
        r.font.color.rgb = GRAY
    return shp


def footer(s, idx: int, appendix: bool = False):
    text(s, "BRM-Net 中文学术汇报" + (" | 备份页" if appendix else ""), 0.62, 7.13, 4.0, 0.18, 9, GRAY)
    text(s, f"{idx:02d}", 12.1, 7.08, 0.55, 0.18, 10, GRAY, align=PP_ALIGN.RIGHT)


def takehome(s, msg: str):
    rect = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.85), Inches(6.25), Inches(11.7), Inches(0.58))
    rect.fill.solid()
    rect.fill.fore_color.rgb = RGBColor(231, 240, 247)
    rect.line.color.rgb = RGBColor(196, 214, 228)
    text(s, msg, 1.05, 6.37, 11.25, 0.25, 21, NAVY, True)


def add_img(s, path: Path, x: float, y: float, w: float, h: float):
    if not path.exists():
        box(s, x, y, w, h, "待补充", "最终实验图表缺失", RED)
        return
    pic = s.shapes.add_picture(str(path), Inches(x), Inches(y), width=Inches(w))
    if pic.height > Inches(h):
        ratio = pic.width / pic.height
        pic.height = Inches(h)
        pic.width = int(pic.height * ratio)
    pic.left = Inches(x + (w - pic.width / 914400) / 2)


def line(s, x1, y1, x2, y2, color=CYAN, dashed=False):
    c = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = color
    c.line.width = Pt(2)
    if dashed:
        c.line.dash_style = 4


def chart(s, cats, series, x, y, w, h, title=None):
    data = ChartData()
    data.categories = cats
    for name, vals in series:
        data.add_series(name, vals)
    frame = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(x), Inches(y), Inches(w), Inches(h), data)
    ch = frame.chart
    ch.has_legend = True
    ch.legend.position = XL_LEGEND_POSITION.BOTTOM
    ch.legend.include_in_layout = False
    ch.value_axis.tick_labels.font.size = Pt(15)
    ch.category_axis.tick_labels.font.size = Pt(15)
    ch.legend.font.size = Pt(15)
    if title:
        ch.has_title = True
        ch.chart_title.text_frame.text = title
        ch.chart_title.text_frame.paragraphs[0].runs[0].font.size = Pt(16)
    return ch


def table(s, headers, rows, x, y, w, h, size=15):
    tbl = s.shapes.add_table(len(rows) + 1, len(headers), Inches(x), Inches(y), Inches(w), Inches(h)).table
    for c, head in enumerate(headers):
        cell = tbl.cell(0, c)
        cell.text = head
        cell.fill.solid()
        cell.fill.fore_color.rgb = NAVY
        p = cell.text_frame.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        p.runs[0].font.name = "Microsoft YaHei"
        p.runs[0].font.size = Pt(size)
        p.runs[0].font.bold = True
        p.runs[0].font.color.rgb = WHITE
    for r, row in enumerate(rows, 1):
        for c, val in enumerate(row):
            cell = tbl.cell(r, c)
            cell.text = val
            cell.fill.solid()
            cell.fill.fore_color.rgb = WHITE if r % 2 else LIGHT
            p = cell.text_frame.paragraphs[0]
            p.alignment = PP_ALIGN.CENTER
            p.runs[0].font.name = "Microsoft YaHei"
            p.runs[0].font.size = Pt(size)
            p.runs[0].font.color.rgb = BLACK


def build(output: Path) -> None:
    formulas = ensure_formulas()
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    s = slide(prs, "BRM-Net：预算化紧凑多模态遥感分类", "开场", dark=True)
    text(s, "面向缺失与退化模态的融合兼容 compact export", 0.78, 1.45, 10.8, 0.45, 24, CYAN, True)
    text(s, "Fusion-Compatible Compact Export with Degradation-Supervised Reliability", 0.78, 2.12, 10.9, 0.35, 18, LIGHT)
    dark_box(s, 0.85, 3.55, 3.25, 1.25, "核心问题", "模型能否真实变小，并稳定处理缺失/退化模态？", CYAN)
    dark_box(s, 4.55, 3.55, 3.25, 1.25, "核心证据", "MAC 对齐、真实导出、鲁棒性", ORANGE)
    dark_box(s, 8.25, 3.55, 3.25, 1.25, "汇报定位", "18 页主汇报\n4 页备份", GREEN)
    text(s, "毕业预答辩 / 学术组会 / 投稿汇报", 0.85, 6.25, 6.0, 0.3, 18, LIGHT)

    s = slide(prs, "研究动机：多模态更强，也更难部署", "01 动机")
    text(s, "HSI", 1.1, 1.75, 1.1, 0.4, 30, BLUE, True)
    text(s, "光谱细节丰富\n但维度高、计算重", 0.85, 2.35, 2.4, 0.8, 21, BLACK)
    text(s, "+", 3.25, 2.2, 0.4, 0.4, 34, ORANGE, True)
    text(s, "LiDAR / Aux", 3.95, 1.68, 2.7, 0.44, 28, GREEN, True)
    text(s, "几何/高度信息\n提升类别区分", 3.95, 2.42, 2.55, 0.82, 21, BLACK)
    text(s, "→", 6.62, 2.18, 0.9, 0.45, 34, CYAN, True, PP_ALIGN.CENTER)
    text(s, "多分支融合", 7.65, 1.75, 2.4, 0.4, 30, NAVY, True)
    text(s, "精度提升\n计算和故障模式增加", 7.65, 2.35, 3.0, 0.8, 21, BLACK)
    flow_card(s, 1.0, 4.08, 10.6, 1.28, "部署导向评价", ["不仅看 full-modality OA", "还报告真实 MAC、参数、缺失和退化状态"], ORANGE)
    takehome(s, "多模态遥感分类需要从“精度优先”转向“精度-效率-鲁棒性联合评价”。")
    footer(s, 2)

    s = slide(prs, "问题：预算、缺失与退化形成三重约束", "02 问题")
    for label, x, y, color in [("资源预算", 5.25, 1.25, CYAN), ("模态缺失", 2.35, 4.15, RED), ("模态退化", 8.15, 4.15, ORANGE)]:
        circ = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(2.0), Inches(1.15))
        circ.fill.solid()
        circ.fill.fore_color.rgb = color
        circ.line.fill.background()
        text(s, label, x + 0.28, y + 0.34, 1.5, 0.28, 22, WHITE, True, PP_ALIGN.CENTER)
    line(s, 6.25, 2.4, 3.35, 4.15, NAVY)
    line(s, 6.25, 2.4, 9.15, 4.15, NAVY)
    line(s, 4.35, 4.72, 8.15, 4.72, NAVY)
    bullets(s, ["导出模型必须真实满足目标资源比例。", "缺失模态需要被排除出融合竞争。", "退化但可用的模态不能简单按缺失处理。"], 0.95, 1.45, 3.2, 1.9, 19)
    bullets(s, ["目标：一个可导出、可解释、可多状态评估的紧凑多模态分类器。"], 8.2, 1.55, 3.8, 0.9, 21)
    takehome(s, "BRM-Net 处理的是三个约束的耦合问题，而不是单独做剪枝或缺失模态补全。")
    footer(s, 3)

    s = slide(prs, "现有方法缺口：四类工作各解决一部分", "03 缺口")
    table(s, ["方法类别", "资源预算", "真实导出", "缺失/退化", "融合兼容"], [
        ["结构剪枝", "强", "强", "弱", "弱"],
        ["弹性网络", "强", "部分", "弱", "弱"],
        ["多模态融合", "弱", "弱", "部分", "强"],
        ["缺失模态学习", "部分", "弱", "强", "强"],
        ["BRM-Net", "强", "强", "强", "强"],
    ], 0.9, 1.55, 11.3, 2.8, 16)
    takehome(s, "本文的空缺定位是：融合接口可保持的真实 compact export + 退化监督可靠融合。")
    footer(s, 4)

    s = slide(prs, "三项核心贡献", "04 贡献")
    text(s, "1", 1.2, 1.55, 0.7, 0.5, 42, CYAN, True)
    flow_card(s, 2.0, 1.22, 9.4, 1.25, "融合兼容预算化紧凑导出", ["学习通道门控", "保持融合接口后导出"], CYAN)
    text(s, "2", 1.2, 3.0, 0.7, 0.5, 42, ORANGE, True)
    flow_card(s, 2.0, 2.73, 9.4, 1.25, "退化监督可靠性融合", ["监督 corrupted-but-available 输入", "不依赖重建网络"], ORANGE)
    text(s, "3", 1.2, 4.45, 0.7, 0.5, 42, GREEN, True)
    flow_card(s, 2.0, 4.24, 9.4, 1.25, "精度-效率-鲁棒性联合评估", ["报告 actual MAC 与参数", "覆盖缺失、退化和多数据集证据"], GREEN)
    takehome(s, "投稿主线收束为三点：真实导出、退化可靠融合、联合评估。")
    footer(s, 5)

    s = slide(prs, "BRM-Net 总体框架：双分支、三阶段", "05 方法")
    text(s, "Training", 1.0, 1.18, 2.0, 0.32, 23, CYAN, True)
    text(s, "Compact Export", 5.25, 1.18, 2.3, 0.32, 23, ORANGE, True)
    text(s, "Inference", 9.65, 1.18, 1.8, 0.32, 23, GREEN, True)

    for y, name, color in [(1.85, "HSI", BLUE), (3.15, "Aux", GREEN)]:
        module_card(s, 0.85, y, 1.25, 0.72, name, color, 20)
        module_card(s, 2.45, y, 1.45, 0.72, "Encoder", color, 19)
        module_card(s, 4.25, y, 1.2, 0.72, "Gate", CYAN, 20)
        line(s, 2.10, y + 0.36, 2.45, y + 0.36)
        line(s, 3.90, y + 0.36, 4.25, y + 0.36)

    module_card(s, 9.15, 1.85, 1.55, 0.72, "Q score", ORANGE, 19)
    module_card(s, 9.15, 3.15, 1.55, 0.72, "Q score", ORANGE, 19)
    module_card(s, 11.0, 2.5, 1.45, 0.78, "Fusion", GREEN, 20)
    module_card(s, 11.0, 3.75, 1.45, 0.78, "Head", NAVY, 20)
    line(s, 5.45, 2.21, 9.15, 2.21)
    line(s, 5.45, 3.51, 9.15, 3.51)
    line(s, 10.70, 2.21, 11.0, 2.72)
    line(s, 10.70, 3.51, 11.0, 3.06)
    line(s, 11.72, 3.28, 11.72, 3.75)

    text(s, "导出路径（训练后执行）", 1.0, 4.58, 2.6, 0.28, 19, ORANGE, True)
    module_card(s, 3.55, 4.88, 1.85, 0.68, "Gate prob.", ORANGE, 18)
    module_card(s, 5.95, 4.88, 1.85, 0.68, "阈值 η", ORANGE, 18)
    module_card(s, 8.35, 4.88, 2.05, 0.68, "通道删除", ORANGE, 18)
    module_card(s, 10.95, 4.88, 1.7, 0.68, "紧凑模型", GREEN, 18)
    line(s, 4.85, 3.90, 4.35, 4.88, ORANGE, True)
    line(s, 5.40, 5.22, 5.95, 5.22, ORANGE, True)
    line(s, 7.80, 5.22, 8.35, 5.22, ORANGE, True)
    line(s, 10.40, 5.22, 10.95, 5.22, ORANGE, True)
    takehome(s, "Compact export 是训练后的结构转换路径，不是普通前向特征流。")
    footer(s, 6)

    s = slide(prs, "预算门控与资源约束", "06 方法")
    add_img(s, formulas["gates"], 0.9, 1.35, 11.0, 0.9)
    add_img(s, formulas["objective"], 0.9, 2.65, 11.0, 0.95)
    chart(s, ["65", "80", "90"], [("目标 MAC", [65, 80, 90]), ("实际 MAC", [64.96, 80.01, 90.07])], 1.1, 4.15, 5.4, 1.8, "Target vs Actual MAC")
    flow_card(s, 7.05, 4.18, 4.85, 1.28, "解释", ["门控用于学习通道重要性", "最终结论基于导出的紧凑模型"], CYAN)
    takehome(s, "实际 MAC 与 65% / 80% / 90% 目标高度一致。")
    footer(s, 7)

    s = slide(prs, "融合兼容 compact export：从 source 到 compact", "07 方法")
    flow_card(s, 0.75, 2.0, 2.65, 1.65, "1. Source model", ["保留原结构", "门控未物理删除"], CYAN)
    flow_card(s, 3.85, 2.0, 2.65, 1.65, "2. Gate projection", ["阈值 η 选择", "终端兼容约束"], ORANGE)
    flow_card(s, 6.95, 2.0, 2.65, 1.65, "3. Physical removal", ["删除通道/BN", "依赖切片"], ORANGE)
    flow_card(s, 10.05, 2.0, 2.45, 1.65, "4. Compact model", ["真实变小", "可部署"], GREEN)
    line(s, 3.40, 2.82, 3.85, 2.82, ORANGE)
    line(s, 6.50, 2.82, 6.95, 2.82, ORANGE)
    line(s, 9.60, 2.82, 10.05, 2.82, GREEN)
    flow_card(s, 1.1, 4.25, 10.9, 1.24, "为什么不是普通剪枝？", ["多模态分支压缩后", "仍需保持融合层终端特征接口"], CYAN)
    takehome(s, "物理导出的紧凑模型保持与 source model 接近的精度。")
    footer(s, 8)

    s = slide(prs, "可用性感知可靠融合：使用论文最终公式", "08 方法")
    add_img(s, formulas["fusion"], 0.65, 1.55, 12.0, 1.05)
    box(s, 1.0, 3.25, 3.3, 1.2, "缺失模态", "可用性为零时\n该模态权重为零", RED, body_size=20)
    box(s, 4.85, 3.25, 3.3, 1.2, "可用模态", "通过质量分数\n和温度进行竞争", CYAN, body_size=20)
    box(s, 8.7, 3.25, 3.3, 1.2, "融合结果", "按可靠性权重\n得到融合特征", GREEN, body_size=20)
    takehome(s, "原简化 softmax 写法已替换为论文中的 masked reliability fusion 公式。")
    footer(s, 9)

    s = slide(prs, "退化监督可靠性：学习 corrupted-but-available 输入", "09 方法")
    add_img(s, formulas["loss"], 0.9, 1.45, 10.8, 0.8)
    for x, title, body, color in [
        (0.95, "干净可用", "质量目标高", GREEN),
        (3.65, "物理缺失", "质量目标为 0", RED),
        (6.35, "噪声/遮挡", "中间质量目标", ORANGE),
        (9.05, "下采样 ×4", "更低质量目标", ORANGE),
    ]:
        box(s, x, 3.05, 2.25, 1.25, title, body, color, body_size=20)
    takehome(s, "可靠性分数是退化监督的融合诊断信号，不宣称物理质量校准。")
    footer(s, 10)

    s = slide(prs, "实验协议：三数据集承担不同角色", "10 实验")
    box(s, 0.9, 1.5, 3.4, 2.1, "Houston2013", "主实验\n官方划分\n预算和消融最完整", CYAN, body_size=20)
    box(s, 4.9, 1.5, 3.4, 2.1, "Trento", "外部验证\n精度接近饱和\n检验效率保持", GREEN, body_size=20)
    box(s, 8.9, 1.5, 3.4, 2.1, "MUUFL", "压力测试\n类别不均衡\n观察 trade-off", ORANGE, body_size=20)
    bullets(s, ["状态：Full、main-only、aux-only、noise、downsample、occlusion。", "指标：OA、AA、Kappa、实际 MAC、参数比例和必要延迟。"], 1.05, 4.35, 10.8, 1.0, 21)
    takehome(s, "三数据集不是排行榜，而是分别支撑主消融、效率保持和困难压力测试。")
    footer(s, 11)

    s = slide(prs, "预算对齐与真实导出结果", "11 结果")
    chart(s, ["65", "80", "90"], [("目标 MAC", [65, 80, 90]), ("实际 MAC", [64.96, 80.01, 90.07]), ("Compact OA", [85.53, 85.84, 85.74])], 0.95, 1.35, 7.1, 3.5, "预算对齐")
    text(s, "64.96%", 8.65, 1.6, 2.4, 0.45, 34, CYAN, True, PP_ALIGN.CENTER)
    text(s, "80.01%", 8.65, 2.65, 2.4, 0.45, 34, ORANGE, True, PP_ALIGN.CENTER)
    text(s, "90.07%", 8.65, 3.7, 2.4, 0.45, 34, GREEN, True, PP_ALIGN.CENTER)
    text(s, "三档目标均稳定导出", 8.35, 4.55, 3.0, 0.3, 20, NAVY, True, PP_ALIGN.CENTER)
    takehome(s, "预算化门控可以被转化为真实 compact model，而不只是训练时软门控。")
    footer(s, 12)

    s = slide(prs, "与压缩 baseline 的 accuracy-efficiency 比较", "12 结果")
    add_img(s, FIG / "fig_pareto_accuracy_efficiency.png", 1.05, 1.25, 7.1, 4.65)
    flow_card(s, 8.65, 1.48, 3.35, 1.42, "读图方式", ["横轴：实际 MAC", "纵轴：OA"], CYAN)
    flow_card(s, 8.65, 3.18, 3.35, 1.42, "关键点", ["80% 目标附近", "比较 compact OA"], ORANGE)
    takehome(s, "该页只保留 Pareto 主图，避免与多面板 overview 重复。")
    footer(s, 13)

    s = slide(prs, "缺失模态鲁棒性：modality dropout 改善 fallback", "13 结果")
    chart(s, ["Full", "HSI only", "LiDAR only", "LiDAR noise"], [
        ("无 dropout", [86.74, 63.34, 32.27, 86.13]),
        ("dropout 0.25", [87.01, 75.84, 42.95, 84.09]),
    ], 0.9, 1.35, 7.2, 3.6, "缺失模态状态")
    text(s, "+12.50 pp", 8.75, 1.75, 2.5, 0.45, 32, CYAN, True, PP_ALIGN.CENTER)
    text(s, "HSI-only", 8.9, 2.25, 2.2, 0.25, 18, GRAY, align=PP_ALIGN.CENTER)
    text(s, "+10.68 pp", 8.75, 3.25, 2.5, 0.45, 32, GREEN, True, PP_ALIGN.CENTER)
    text(s, "LiDAR-only", 8.9, 3.75, 2.2, 0.25, 18, GRAY, align=PP_ALIGN.CENTER)
    takehome(s, "训练时暴露缺失状态显著提升单模态 fallback。")
    footer(s, 14)

    s = slide(prs, "连续退化鲁棒性：adverse average 明显提升", "14 结果")
    chart(s, ["Baseline", "Uniform", "Noise", "Multi .50", "Multi .25"], [
        ("Adverse Avg.", [74.89, 76.41, 79.54, 84.72, 84.76]),
        ("Full OA", [86.97, 86.72, 87.45, 85.99, 86.90]),
    ], 0.85, 1.35, 7.6, 3.6, "Clean vs degraded states")
    text(s, "74.89 → 84.76", 8.75, 1.95, 3.25, 0.42, 28, ORANGE, True, PP_ALIGN.CENTER)
    text(s, "adverse-state average", 8.75, 2.55, 3.25, 0.25, 18, GRAY, align=PP_ALIGN.CENTER)
    text(s, "Full OA 保持 86.90", 8.65, 3.45, 3.3, 0.35, 22, NAVY, True, PP_ALIGN.CENTER)
    takehome(s, "multi-degradation p=0.25 在鲁棒性与 clean OA 之间最平衡。")
    footer(s, 15)

    s = slide(prs, "多数据集验证：不同数据集回答不同问题", "15 结果")
    table(s, ["数据集", "Full OA", "Main-only", "Occlusion-50", "MACs"], [
        ["Houston2013", "86.90±1.65", "80.04±1.93", "85.57±1.64", "80.01"],
        ["Trento", "98.94±0.94", "92.89±3.15", "98.16±1.87", "79.90"],
        ["MUUFL", "86.87±1.79", "83.03±2.14", "85.99±2.08", "79.97"],
    ], 0.8, 1.45, 11.8, 2.2, 16)
    box(s, 1.0, 4.25, 3.2, 1.0, "Houston", "主消融证据", CYAN, body_size=21)
    box(s, 5.05, 4.25, 3.2, 1.0, "Trento", "效率保持证据", GREEN, body_size=21)
    box(s, 9.1, 4.25, 3.2, 1.0, "MUUFL", "困难 trade-off 证据", ORANGE, body_size=21)
    takehome(s, "不要把三数据集写成简单排行榜，而要写成角色分工证据。")
    footer(s, 16)

    s = slide(prs, "消融、trade-off 与局限", "16 讨论")
    table(s, ["项目", "说明", "结论"], [
        ["w/o mask", "去掉可用性掩码", "HSI-only 下降 3.32 pp"],
        ["Noise quality", "只监督噪声", "遮挡改善不足"],
        ["Multi-deg .25", "轻量多退化", "adverse average 最优"],
    ], 0.9, 1.35, 11.2, 1.8, 16)
    flow_card(s, 0.95, 3.68, 3.45, 1.42, "已支撑", ["预算对齐、真实导出", "退化鲁棒性"], GREEN)
    flow_card(s, 4.85, 3.68, 3.45, 1.42, "需谨慎", ["MUUFL 存在", "clean-robustness trade-off"], ORANGE)
    flow_card(s, 8.75, 3.68, 3.45, 1.42, "下一步", ["更大 backbone 与更多模态", "更密集退化曲线"], CYAN)
    takehome(s, "本文主张是联合权衡，不是所有场景 full OA 都单调提升。")
    footer(s, 17)

    s = slide(prs, "总结与 Q&A", "17 总结", dark=True)
    text(s, "一句话结论", 0.85, 1.1, 2.6, 0.42, 28, CYAN, True)
    text(s, "BRM-Net 面向资源受限场景，学习可真实导出的紧凑多模态分类器，并显式评估缺失与退化模态。", 0.85, 1.85, 11.2, 0.9, 30, WHITE, True)
    dark_box(s, 0.95, 3.75, 3.45, 1.25, "贡献 1", "融合兼容预算化紧凑导出", CYAN)
    dark_box(s, 4.85, 3.75, 3.45, 1.25, "贡献 2", "退化监督可靠性融合", ORANGE)
    dark_box(s, 8.75, 3.75, 3.45, 1.25, "贡献 3", "精度-效率-鲁棒性联合评估", GREEN)
    text(s, "Q&A", 0.95, 6.25, 2.5, 0.45, 34, CYAN, True)

    s = slide(prs, "备份 1：retained width 结构诊断", "备份")
    add_img(s, FIG / "fig_gate_retention.png", 2.0, 1.2, 8.9, 4.9)
    takehome(s, "保留宽度用于解释 learned nonuniform export 的层级分配。")
    footer(s, 19, True)

    s = slide(prs, "备份 2：退化可靠性诊断曲线", "备份")
    add_img(s, FIG / "fig_controlled_corruption_reliability_curves.png", 1.4, 1.1, 10.2, 5.0)
    takehome(s, "可靠性曲线是机制诊断，不是物理质量校准声明。")
    footer(s, 20, True)

    s = slide(prs, "备份 3：MUUFL 错误分析", "备份")
    add_img(s, FIG / "fig_muufl_error_analysis.png", 1.3, 1.2, 10.8, 4.8)
    takehome(s, "MUUFL 用于说明鲁棒训练可能带来类别相关的 clean-robustness trade-off。")
    footer(s, 21, True)

    s = slide(prs, "备份 4：毕业论文扩展方向", "备份", dark=True)
    dark_box(s, 0.9, 1.55, 3.45, 1.45, "动态路由", "面向不同预算和模态状态选择 profile。", CYAN)
    dark_box(s, 4.9, 1.55, 3.45, 1.45, "加权交叉熵", "处理 MUUFL 等类别不均衡问题。", ORANGE)
    dark_box(s, 8.9, 1.55, 3.45, 1.45, "演示系统", "交互选择数据集、预算和模态状态。", GREEN)
    text(s, "这些内容不进入投稿主线，作为毕业论文工作量和答辩演示补充。", 1.0, 4.9, 11.0, 0.45, 26, WHITE, True)

    output.parent.mkdir(parents=True, exist_ok=True)
    prs.save(output)


def main() -> int:
    args = parser().parse_args()
    build(Path(args.output))
    print(f"Wrote {args.output}")
    print(f"Slides: {len(SLIDES)}")
    for i, t in enumerate(SLIDES, 1):
        print(f"{i:02d}. {t}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
