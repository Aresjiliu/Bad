from __future__ import annotations

import argparse
import sys
from pathlib import Path


ROOT = Path(__file__).resolve().parents[1]
PPT_DEPS = ROOT / "output" / "ppt_deps"
if PPT_DEPS.exists():
    sys.path.insert(0, str(PPT_DEPS))

from pptx import Presentation  # type: ignore  # noqa: E402
from pptx.dml.color import RGBColor  # type: ignore  # noqa: E402
from pptx.enum.text import PP_ALIGN  # type: ignore  # noqa: E402
from pptx.util import Inches  # type: ignore  # noqa: E402

from make_brmnet_academic_presentation import (  # noqa: E402
    BLACK,
    CYAN,
    CYAN_DARK,
    FIG,
    GREEN,
    GRAY,
    LIGHT,
    NAVY,
    NAVY2,
    OFFWHITE,
    ORANGE,
    RED,
    WHITE,
    add_picture,
    bullet_list,
    card,
    chart_bar,
    connector,
    dark_card,
    kpi,
    multiline,
    simple_table,
    textbox,
)


OUT = ROOT / "docs" / "slides" / "BRM-Net_Academic_Presentation_ZH.pptx"
SLIDE_TITLES: list[str] = []


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description="Build the Chinese BRM-Net academic presentation.")
    parser.add_argument("--output", default=str(OUT))
    return parser


def zh_textbox(slide, *args, **kwargs):
    kwargs.setdefault("font", "Microsoft YaHei")
    return textbox(slide, *args, **kwargs)


def zh_slide(prs: Presentation, title: str, section: str, dark: bool = False):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = NAVY if dark else OFFWHITE
    if not dark:
        bar = slide.shapes.add_shape(1, Inches(0), Inches(0), Inches(13.333), Inches(0.18))
        bar.fill.solid()
        bar.fill.fore_color.rgb = NAVY
        bar.line.fill.background()
        zh_textbox(slide, section.upper(), 0.63, 0.31, 2.2, 0.24, size=8, color=CYAN_DARK, bold=True)
        zh_textbox(slide, title, 0.63, 0.53, 11.4, 0.52, size=27, color=NAVY, bold=True)
    SLIDE_TITLES.append(title)
    return slide


def zh_footer(slide, idx: int):
    zh_textbox(slide, "BRM-Net 中文学术汇报", 0.63, 7.12, 4.2, 0.18, size=8, color=GRAY)
    zh_textbox(slide, f"{idx:02d}", 12.15, 7.06, 0.48, 0.2, size=10, color=GRAY, align=PP_ALIGN.RIGHT)


def build(output: Path) -> None:
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    s = zh_slide(prs, "标题页", "开场", dark=True)
    zh_textbox(s, "BRM-Net", 0.72, 0.72, 4.2, 0.45, size=21, color=CYAN, bold=True)
    zh_textbox(s, "面向多模态遥感分类的\n融合兼容预算化紧凑导出与退化监督可靠性学习", 0.72, 1.42, 11.6, 1.25, size=30, color=WHITE, bold=True)
    zh_textbox(s, "毕业预答辩 / 学术组会 / 投稿汇报中文版本", 0.76, 2.92, 8.5, 0.35, size=17, color=LIGHT)
    dark_card(s, 0.78, 4.1, 3.4, 1.2, "核心问题", "多模态模型能否既真实变小，又能应对模态缺失与退化？", CYAN)
    dark_card(s, 4.6, 4.1, 3.4, 1.2, "主要证据", "目标 MAC 对齐、真实 compact export、多状态鲁棒性评估。", ORANGE)
    dark_card(s, 8.42, 4.1, 3.4, 1.2, "汇报边界", "聚焦投稿主线，毕业论文扩展放在未来工作。", GREEN)
    zh_textbox(s, "BRM-Net | Fusion-Compatible Budgeted Compact Export", 0.78, 6.62, 7.2, 0.28, size=11, color=LIGHT)

    s = zh_slide(prs, "研究动机：为什么需要轻量化多模态遥感分类", "动机")
    kpi(s, 0.95, 1.55, 2.1, "2+", "多传感器分支", CYAN)
    kpi(s, 3.55, 1.55, 2.1, "80%", "典型目标 MAC", ORANGE)
    kpi(s, 6.15, 1.55, 2.1, "3", "数据集证据", GREEN)
    card(s, 8.95, 1.18, 3.2, 2.35, "部署压力", "边缘端、星载端和资源受限平台需要更低计算、更低参数和更稳定的推理行为。", ORANGE)
    bullet_list(s, [
        "HSI 提供丰富光谱信息，但维度高、计算成本大。",
        "LiDAR 等辅助模态提供几何或空间信息，能提升识别。",
        "多分支融合提高精度，但也放大计算开销和故障模式。",
        "最终模型必须报告真实 MAC/参数下降，而不是只看训练时软门控。",
    ], 0.95, 4.05, 10.9, 1.65, size=16)
    zh_footer(s, 2)

    s = zh_slide(prs, "问题定义：资源预算、模态缺失与模态退化耦合存在", "问题")
    card(s, 0.8, 1.35, 3.3, 2.0, "资源预算", "导出的最终网络需要满足给定资源比例 rho，并以实际 MAC/参数验证。", CYAN)
    card(s, 4.35, 1.35, 3.3, 2.0, "模态缺失", "某个模态可能物理不可用，需要从融合 softmax 竞争中排除。", RED)
    card(s, 7.9, 1.35, 3.3, 2.0, "模态退化", "模态仍然存在，但可能受到噪声、低分辨率或遮挡影响。", ORANGE)
    zh_textbox(s, "期望模型行为", 0.9, 4.25, 3.2, 0.3, size=19, color=NAVY, bold=True)
    bullet_list(s, [
        "完整模态精度不能明显崩塌。",
        "导出后模型确实具有更少通道和更低计算量。",
        "能覆盖 full、main-only、aux-only、corrupted auxiliary 等推理状态。",
        "可靠性分数只作为退化监督诊断信号，不宣称物理校准。",
    ], 0.95, 4.72, 10.5, 1.45, size=16)
    zh_footer(s, 3)

    s = zh_slide(prs, "现有工作的关键空缺", "差距")
    simple_table(s, ["方法类别", "预算", "真实导出", "缺失模态", "退化模态", "融合接口"], [
        ["结构剪枝", "有", "有", "无", "无", "多为单主干"],
        ["弹性网络", "有", "部分", "无", "无", "宽度自适应"],
        ["多模态遥感融合", "少", "无", "有限", "有限", "精度优先"],
        ["缺失模态学习", "有限", "无", "有", "部分", "常依赖重模块"],
        ["BRM-Net", "有", "有", "有", "有", "显式保持"],
    ], 0.65, 1.45, 12.0, 2.5, font_size=8)
    card(s, 0.9, 4.55, 3.5, 1.15, "差距 1", "普通剪枝很少考虑多模态融合接口约束。", CYAN)
    card(s, 4.9, 4.55, 3.5, 1.15, "差距 2", "缺失模态方法常依赖重建或专家系统，与轻量化目标冲突。", ORANGE)
    card(s, 8.9, 4.55, 3.5, 1.15, "差距 3", "退化但仍可用的辅助模态没有被充分测试。", RED)
    zh_footer(s, 4)

    s = zh_slide(prs, "BRM-Net 总体框架", "方法")
    xs = [0.7, 3.2, 5.7, 8.2, 10.7]
    titles = ["输入", "门控编码器", "可靠性分支", "掩码融合", "紧凑导出"]
    bodies = ["HSI/main\n辅助模态", "Hard-concrete\n通道门控", "质量分数\nq_m", "可用性约束\nsoftmax", "物理删除\n冗余通道"]
    colors = [NAVY, CYAN, ORANGE, GREEN, NAVY2]
    for i, x in enumerate(xs):
        card(s, x, 2.0, 2.0, 1.65, titles[i], bodies[i], colors[i])
        if i < len(xs) - 1:
            connector(s, x + 2.0, 2.83, x + 2.45, 2.83, CYAN)
    zh_textbox(s, "训练阶段", 1.4, 4.35, 1.3, 0.28, size=15, color=CYAN_DARK, bold=True)
    connector(s, 2.45, 4.48, 5.15, 4.48, CYAN)
    zh_textbox(s, "紧凑导出", 5.25, 4.35, 2.0, 0.28, size=15, color=ORANGE, bold=True)
    connector(s, 7.0, 4.48, 9.25, 4.48, ORANGE)
    zh_textbox(s, "部署推理", 9.35, 4.35, 1.5, 0.28, size=15, color=GREEN, bold=True)
    zh_textbox(s, "核心原则：压缩内部通道，但保持融合层所需的终端特征接口。", 1.1, 5.7, 10.8, 0.4, size=18, color=NAVY, bold=True, align=PP_ALIGN.CENTER)
    zh_footer(s, 5)

    s = zh_slide(prs, "问题建模：从训练约束到可部署 compact model", "方法")
    zh_textbox(s, "约束优化目标", 0.9, 1.38, 3.3, 0.36, size=20, color=NAVY, bold=True)
    card(s, 0.92, 1.9, 4.8, 1.7, "部署目标", "min E[L_cls(f_theta(g(x), a), y)]\n约束：R(f_theta^c) <= rho R(f_theta)", CYAN)
    multiline(s, [
        ("x^(1), x^(2)：", "异构对齐观测，如 HSI-LiDAR。"),
        ("a_m：", "可用性标记；缺失模态从融合中屏蔽。"),
        ("g：", "退化算子，包括噪声、下采样、遮挡。"),
        ("f_theta^c：", "结构投影后的物理紧凑模型。"),
    ], 6.25, 1.42, 5.7, 2.35, size=14)
    card(s, 1.1, 4.55, 10.9, 1.25, "谨慎表述", "本文不声称可靠性分数是物理传感器质量的校准值，而是将其作为退化监督下的融合诊断信号。", ORANGE)
    zh_footer(s, 6)

    s = zh_slide(prs, "预算感知通道门控", "方法")
    card(s, 0.85, 1.35, 3.35, 2.1, "门控定义", "每个被门控层具有 z_l in [0,1]^C，对特征通道进行逐通道缩放。", CYAN)
    card(s, 4.75, 1.35, 3.35, 2.1, "训练估计", "Hard-concrete 概率用于估计期望通道使用量和资源消耗。", ORANGE)
    card(s, 8.65, 1.35, 3.35, 2.1, "导出规则", "选择阈值并二值化门控，使实际 MAC 接近目标预算。", GREEN)
    chart_bar(s, ["65", "80", "90"], [("目标", [65, 80, 90]), ("实际", [64.96, 80.01, 90.07])], 1.2, 4.1, 5.2, 2.2, "MAC 预算对齐")
    bullet_list(s, [
        "资源损失约束期望 MAC 与目标预算之间的偏差。",
        "门控只是训练中间机制。",
        "论文最终报告的是导出后的 compact model。",
    ], 7.0, 4.15, 4.8, 1.55, size=15)
    zh_footer(s, 7)

    s = zh_slide(prs, "融合兼容紧凑导出", "方法")
    card(s, 0.85, 1.35, 3.3, 1.65, "为什么需要兼容", "独立剪掉分支通道可能破坏融合头所期望的特征维度。", RED)
    card(s, 4.55, 1.35, 3.3, 1.65, "结构投影", "删除输出通道、下一层输入切片、BN 参数和相关分类器权重。", CYAN)
    card(s, 8.25, 1.35, 3.3, 1.65, "终端约束", "通过保护或绑定终端门控，保持融合使用的表示维度。", GREEN)
    simple_table(s, ["状态", "Source OA", "Compact OA", "变化", "MACs"], [
        ["Full", "86.15", "86.90", "+0.75", "80.01"],
        ["HSI only", "78.18", "80.04", "+1.86", "80.01"],
        ["Noise-high", "84.26", "84.55", "+0.29", "80.01"],
        ["Occlusion-50", "84.59", "85.57", "+0.98", "80.01"],
    ], 1.0, 4.0, 10.9, 1.65, font_size=9)
    zh_textbox(s, "Source gated model 与真实导出的 compact model 表现接近，支撑真实导出主张。", 1.05, 6.1, 10.8, 0.3, size=14, color=NAVY, bold=True)
    zh_footer(s, 8)

    s = zh_slide(prs, "可用性感知可靠性融合", "方法")
    card(s, 0.85, 1.35, 3.45, 2.0, "可靠性分数", "每个模态输出 q_m，作为融合权重的质量信号。", ORANGE)
    card(s, 4.72, 1.35, 3.45, 2.0, "可用性掩码", "不可用模态在 softmax 前被赋予极小 logit。", CYAN)
    card(s, 8.59, 1.35, 3.45, 2.0, "融合权重", "alpha_m = softmax(log q_m + mask(a_m))，仅在可用模态间竞争。", GREEN)
    simple_table(s, ["状态", "BRM-Net", "去掉 mask", "变化"], [
        ["Full", "86.90", "87.02", "+0.12"],
        ["HSI only", "80.04", "76.72", "-3.32"],
        ["LiDAR only", "38.04", "37.59", "-0.45"],
        ["Occlusion-50", "85.57", "85.67", "+0.10"],
    ], 1.2, 4.15, 10.2, 1.55, font_size=10)
    zh_textbox(s, "结论：mask 主要保护物理缺失模态，而不是所有退化但可用的情况。", 1.15, 6.1, 10.6, 0.28, size=14, color=NAVY, bold=True)
    zh_footer(s, 9)

    s = zh_slide(prs, "退化监督质量学习", "方法")
    card(s, 0.8, 1.25, 2.65, 1.5, "干净且可用", "质量目标 t_m = 1", GREEN)
    card(s, 3.75, 1.25, 2.65, 1.5, "物理缺失", "质量目标 t_m = 0", RED)
    card(s, 6.7, 1.25, 2.65, 1.5, "噪声 / 遮挡", "中间质量目标", ORANGE)
    card(s, 9.65, 1.25, 2.65, 1.5, "下采样 x4", "更低质量目标", ORANGE)
    zh_textbox(s, "训练损失", 0.95, 3.75, 2.0, 0.3, size=19, color=NAVY, bold=True)
    card(s, 0.95, 4.2, 5.05, 1.05, "目标函数", "L = L_cls + lambda_b L_bud + lambda_q L_qual", CYAN)
    bullet_list(s, [
        "退化模态仍参与候选融合，而不是直接屏蔽。",
        "可靠性分支需要学会降低受损辅助模态权重。",
        "该分数是退化感知诊断信号，不是物理校准量。",
    ], 6.65, 3.85, 5.2, 1.45, size=15)
    zh_footer(s, 10)

    s = zh_slide(prs, "训练与导出流程", "方法")
    phases = [
        ("1", "源模型训练", "模态 dropout + 退化采样 + 预算损失"),
        ("2", "门控投影", "选择阈值 eta 匹配目标 MAC"),
        ("3", "紧凑导出", "删除通道和依赖参数"),
        ("4", "短程微调", "恢复 compact 架构精度"),
        ("5", "多状态评估", "Full、缺失、退化、MAC、参数、延迟"),
    ]
    for i, (num, title, body) in enumerate(phases):
        x = 0.75 + i * 2.48
        card(s, x, 1.65, 2.08, 2.15, f"{num}. {title}", body, [CYAN, ORANGE, GREEN, NAVY2, CYAN_DARK][i])
        if i < len(phases) - 1:
            connector(s, x + 2.08, 2.72, x + 2.42, 2.72, CYAN)
    card(s, 1.15, 4.85, 10.9, 1.1, "算法重点", "导出阶段不可省略：部署模型具有物理更少通道，而不是仍保留原网络并加软门控。", ORANGE)
    zh_footer(s, 11)

    s = zh_slide(prs, "实验协议与数据集", "实验")
    simple_table(s, ["数据集", "角色", "训练 / 测试", "纳入原因"], [
        ["Houston2013", "主实验", "2,832 / 12,197", "官方划分，消融最完整"],
        ["Trento", "外部验证", "819 / 29,395", "精度接近饱和的 HSI-LiDAR 场景"],
        ["MUUFL", "压力测试", "1,550 / 52,137", "类别不均衡，辅助模态较弱"],
    ], 0.75, 1.4, 11.9, 2.0, font_size=9)
    card(s, 0.9, 4.2, 3.35, 1.35, "指标", "OA、AA、Kappa、MAC ratio、参数 ratio、必要时报告 latency。", CYAN)
    card(s, 4.95, 4.2, 3.35, 1.35, "状态", "Full、main-only、aux-only、noise、downsample、occlusion。", ORANGE)
    card(s, 9.0, 4.2, 3.35, 1.35, "汇报原则", "三数据集承担不同角色，不写成单一排行榜。", GREEN)
    zh_footer(s, 12)

    s = zh_slide(prs, "预算对齐结果", "结果")
    kpi(s, 1.05, 1.4, 2.0, "64.96%", "65 目标下实际 MAC", CYAN)
    kpi(s, 4.0, 1.4, 2.0, "80.01%", "80 目标下实际 MAC", ORANGE)
    kpi(s, 6.95, 1.4, 2.0, "90.07%", "90 目标下实际 MAC", GREEN)
    chart_bar(s, ["65", "80", "90"], [("实际 MAC", [64.96, 80.01, 90.07]), ("Compact OA", [85.53, 85.84, 85.74])], 1.15, 3.5, 5.7, 2.45, "预算与精度")
    simple_table(s, ["目标", "Source OA", "Compact OA", "Params", "MACs"], [
        ["65", "85.98", "85.53", "64.79", "64.96"],
        ["80", "85.79", "85.84", "79.92", "80.01"],
        ["90", "85.33", "85.74", "89.78", "90.07"],
    ], 7.25, 3.55, 4.55, 1.75, font_size=9)
    zh_textbox(s, "主张：目标资源预算可以被物理导出为实际 compact network。", 1.0, 6.28, 10.8, 0.3, size=15, color=NAVY, bold=True)
    zh_footer(s, 13)

    s = zh_slide(prs, "精度-效率权衡", "结果")
    add_picture(s, FIG / "fig_pareto_accuracy_efficiency.png", 0.85, 1.25, 5.25, 4.75)
    add_picture(s, FIG / "fig_brmnet_results_overview.png", 6.35, 1.25, 5.75, 4.75)
    zh_textbox(s, "左图突出 accuracy-efficiency trade-off，右图汇总缺失模态、保留宽度和预算对齐证据。", 0.95, 6.18, 10.8, 0.32, size=14, color=NAVY, bold=True)
    zh_footer(s, 14)

    s = zh_slide(prs, "模态缺失鲁棒性", "结果")
    simple_table(s, ["训练方式", "Full", "HSI only", "LiDAR only", "LiDAR noise"], [
        ["无 dropout", "86.74", "63.34", "32.27", "86.13"],
        ["dropout 0.25", "87.01", "75.84", "42.95", "84.09"],
    ], 0.95, 1.35, 5.6, 1.15, font_size=10)
    chart_bar(s, ["Full", "HSI only", "LiDAR only", "Noise"], [
        ("无 dropout", [86.74, 63.34, 32.27, 86.13]),
        ("dropout 0.25", [87.01, 75.84, 42.95, 84.09]),
    ], 0.95, 3.0, 5.9, 2.75, "Modality dropout 效果")
    card(s, 7.35, 1.4, 4.2, 1.35, "观察", "训练时模态 dropout 显著提升单模态 fallback 能力。", GREEN)
    card(s, 7.35, 3.15, 4.2, 1.35, "边界", "缺失由 availability mask 处理，退化仍需要质量监督。", ORANGE)
    zh_footer(s, 15)

    s = zh_slide(prs, "模态退化鲁棒性", "结果")
    simple_table(s, ["方法", "Full", "Noise-high", "Down-4", "Occ-50", "Adverse"], [
        ["Full baseline", "86.97", "72.08", "78.37", "74.21", "74.89"],
        ["Uniform fusion", "86.72", "77.01", "78.35", "73.87", "76.41"],
        ["Noise quality", "87.45", "85.15", "80.14", "73.32", "79.54"],
        ["Multi-deg .50", "85.99", "85.11", "84.38", "84.66", "84.72"],
        ["Multi-deg .25", "86.90", "84.55", "84.15", "85.57", "84.76"],
    ], 0.65, 1.28, 12.0, 2.15, font_size=8)
    chart_bar(s, ["Baseline", "Uniform", "Noise", "Multi .50", "Multi .25"], [
        ("Adverse Avg.", [74.89, 76.41, 79.54, 84.72, 84.76]),
        ("Full OA", [86.97, 86.72, 87.45, 85.99, 86.90]),
    ], 1.15, 4.0, 6.1, 2.35, "Clean 与 adverse 状态")
    card(s, 7.85, 4.2, 3.9, 1.3, "关键结果", "80% MAC 设定下，adverse-state average 从 74.89 提升到 84.76。", ORANGE)
    zh_footer(s, 16)

    s = zh_slide(prs, "多数据集证据：Houston / Trento / MUUFL", "结果")
    simple_table(s, ["数据集", "Full OA", "Main-only", "Occlusion-50", "MACs"], [
        ["Houston2013", "86.90 ± 1.65", "80.04 ± 1.93", "85.57 ± 1.64", "80.01"],
        ["Trento", "98.94 ± 0.94", "92.89 ± 3.15", "98.16 ± 1.87", "79.90"],
        ["MUUFL", "86.87 ± 1.79", "83.03 ± 2.14", "85.99 ± 2.08", "79.97"],
    ], 0.75, 1.3, 11.75, 1.65, font_size=10)
    chart_bar(s, ["Houston", "Trento", "MUUFL"], [
        ("Full OA", [86.90, 98.94, 86.87]),
        ("Main-only", [80.04, 92.89, 83.03]),
        ("Occlusion-50", [85.57, 98.16, 85.99]),
    ], 1.0, 3.55, 6.0, 2.65, "跨数据集角色证据")
    card(s, 7.65, 3.7, 3.9, 1.65, "解释", "Houston 支撑主消融；Trento 证明轻量化不破坏近饱和精度；MUUFL 暴露 clean-robustness 权衡。", CYAN)
    zh_footer(s, 17)

    s = zh_slide(prs, "消融实验结论", "结果")
    simple_table(s, ["消融项", "改变内容", "观察结果"], [
        ["Uniform fusion", "去掉可靠性加权", "退化鲁棒性有限"],
        ["Noise quality", "只监督噪声质量", "噪声改善，遮挡仍弱"],
        ["Multi-deg .50", "更强退化日程", "鲁棒但 full OA 略降"],
        ["Multi-deg .25", "轻量多退化日程", "当前 adverse average 最优"],
        ["w/o mask", "去掉可用性掩码", "HSI-only OA 下降 3.32 pp"],
    ], 0.75, 1.35, 11.85, 2.4, font_size=9)
    card(s, 0.9, 4.55, 5.0, 1.35, "结论", "鲁棒性主要来自退化监督可靠性，而不是单纯加入一个可靠性标量。", ORANGE)
    card(s, 6.65, 4.55, 5.0, 1.35, "投稿边界", "routing、weighted CE、profile label stabilization 放入毕业论文扩展。", CYAN)
    zh_footer(s, 18)

    s = zh_slide(prs, "可视化：保留宽度、可靠性诊断与 MUUFL 错误分析", "可视化")
    add_picture(s, FIG / "fig_gate_retention.png", 0.75, 1.25, 3.35, 4.55)
    add_picture(s, FIG / "fig_controlled_corruption_reliability_curves.png", 4.35, 1.25, 4.35, 4.55)
    add_picture(s, FIG / "fig_muufl_error_analysis.png", 8.95, 1.25, 3.45, 4.55)
    zh_textbox(s, "这些图用于机制诊断：可靠性是 degradation-aware，而不是 physically calibrated。", 0.95, 6.22, 10.8, 0.3, size=14, color=NAVY, bold=True)
    zh_footer(s, 19)

    s = zh_slide(prs, "讨论：clean-robustness 权衡与局限", "讨论")
    card(s, 0.85, 1.35, 3.4, 1.65, "已支撑", "Houston2013 上目标 MAC 对齐和 compact export 证据较完整。", GREEN)
    card(s, 4.8, 1.35, 3.4, 1.65, "需谨慎", "MUUFL 增强鲁棒性，但 clean full OA 牺牲 1.64 pp。", ORANGE)
    card(s, 8.75, 1.35, 3.4, 1.65, "局限", "当前多数据集仍以 HSI-LiDAR 为主，HS-MS 与 HS-SAR 仍需扩展。", RED)
    bullet_list(s, [
        "单一 full-modality OA 不能充分评价部署导向模型。",
        "可靠性加权必须配合 corrupted-but-available 训练状态才有效。",
        "不能夸大为真实星载部署或物理传感器质量校准。",
        "更大 backbone 和更强 lightweight baseline 会进一步增强说服力。",
    ], 1.0, 4.0, 10.8, 1.65, size=16)
    zh_footer(s, 20)

    s = zh_slide(prs, "贡献总结", "总结")
    kpi(s, 0.95, 1.45, 2.0, "1", "紧凑导出", CYAN)
    kpi(s, 3.6, 1.45, 2.0, "2", "可靠融合", ORANGE)
    kpi(s, 6.25, 1.45, 2.0, "3", "联合评估", GREEN)
    card(s, 0.95, 3.55, 3.4, 1.6, "融合兼容预算化紧凑导出", "学习门控并真实导出 compact model，同时保持融合接口。", CYAN)
    card(s, 4.85, 3.55, 3.4, 1.6, "退化监督可靠性融合", "为 corrupted-but-available 模态提供质量目标，不依赖重建网络。", ORANGE)
    card(s, 8.75, 3.55, 3.4, 1.6, "精度-效率-鲁棒性联合评估", "报告实际 MAC、缺失状态、退化状态和多数据集证据。", GREEN)
    zh_footer(s, 21)

    s = zh_slide(prs, "未来工作与 Q&A", "总结", dark=True)
    zh_textbox(s, "未来工作", 0.75, 0.72, 3.5, 0.45, size=28, color=WHITE, bold=True)
    dark_card(s, 0.85, 1.65, 3.45, 1.45, "毕业论文扩展", "动态路由、weighted CE、profile label stabilization、类别级诊断。", CYAN)
    dark_card(s, 4.85, 1.65, 3.45, 1.45, "增强证据", "更大 backbone、更强轻量化 baseline、更密集退化曲线。", ORANGE)
    dark_card(s, 8.85, 1.65, 3.45, 1.45, "演示系统", "交互式选择数据集、模态状态和预算配置，服务答辩展示。", GREEN)
    zh_textbox(s, "Q&A", 0.85, 4.45, 3.0, 0.55, size=34, color=CYAN, bold=True)
    zh_textbox(s, "一句话总结：BRM-Net 关注资源受限场景下，可真实导出的紧凑多模态遥感分类，并显式评估缺失与退化模态。", 0.9, 5.25, 10.8, 0.6, size=18, color=WHITE, bold=True)
    zh_textbox(s, "待补充：最终混淆矩阵、最终延迟表、更密集退化曲线。", 0.9, 6.42, 10.8, 0.3, size=13, color=LIGHT)

    output.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(output))


def main() -> int:
    args = build_parser().parse_args()
    output = Path(args.output)
    build(output)
    print(f"Wrote {output}")
    print(f"Slides: {len(SLIDE_TITLES)}")
    for i, title in enumerate(SLIDE_TITLES, 1):
        print(f"{i:02d}. {title}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
