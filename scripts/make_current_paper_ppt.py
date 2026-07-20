from __future__ import annotations

import argparse
import sys
from pathlib import Path


PPT_DEPS = Path(__file__).resolve().parents[1] / "output" / "ppt_deps"
if PPT_DEPS.exists():
    sys.path.insert(0, str(PPT_DEPS))

from pptx import Presentation  # type: ignore  # noqa: E402
from pptx.dml.color import RGBColor  # type: ignore  # noqa: E402
from pptx.enum.shapes import MSO_SHAPE  # type: ignore  # noqa: E402
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR  # type: ignore  # noqa: E402
from pptx.util import Inches, Pt  # type: ignore  # noqa: E402


ROOT = Path(__file__).resolve().parents[1]
DEFAULT_TEMPLATE = Path("D:/Academic/ppt/Elegant Black & White Thesis Defense by Slidesgo.pptx")
DEFAULT_OUTPUT = ROOT / "docs" / "slides" / "BRMNet_submission_report_20260720.pptx"
PAPER_DIR = Path("D:/Academic/paper_submission/brmnet_pricai2026")


BLACK = RGBColor(18, 18, 18)
WHITE = RGBColor(255, 255, 255)
OFFWHITE = RGBColor(247, 247, 244)
GRAY = RGBColor(100, 100, 96)
LIGHT = RGBColor(232, 232, 226)
ACCENT = RGBColor(177, 143, 78)
GREEN = RGBColor(66, 137, 108)
BLUE = RGBColor(59, 104, 154)
RED = RGBColor(170, 82, 75)


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description="Build the current BRM-Net paper report deck.")
    parser.add_argument("--template", default=str(DEFAULT_TEMPLATE))
    parser.add_argument("--output", default=str(DEFAULT_OUTPUT))
    return parser


def clear_template_slides(prs: Presentation) -> None:
    slide_id_list = prs.slides._sldIdLst  # noqa: SLF001
    for slide_id in list(slide_id_list):
        r_id = slide_id.rId
        prs.part.drop_rel(r_id)
        slide_id_list.remove(slide_id)


def add_slide(prs: Presentation, dark: bool = False):
    slide = prs.slides.add_slide(prs.slide_layouts[10])
    bg = slide.background.fill
    bg.solid()
    bg.fore_color.rgb = BLACK if dark else OFFWHITE
    return slide


def add_text(
    slide,
    text: str,
    x: float,
    y: float,
    w: float,
    h: float,
    size: int = 18,
    bold: bool = False,
    color=BLACK,
    align=PP_ALIGN.LEFT,
    font: str = "Microsoft YaHei",
):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.TOP
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    return box


def add_title(slide, title: str, subtitle: str | None = None, dark: bool = False):
    color = WHITE if dark else BLACK
    add_text(slide, title, 0.65, 0.38, 11.0, 0.7, size=28, bold=True, color=color)
    if subtitle:
        add_text(slide, subtitle, 0.68, 1.08, 10.8, 0.35, size=11, color=ACCENT if dark else GRAY)


def add_rule(slide, x: float, y: float, w: float, color=ACCENT, height: float = 0.03):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(height))
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()
    return shape


def add_card(slide, x: float, y: float, w: float, h: float, title: str, body: str, accent=ACCENT):
    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    card.fill.solid()
    card.fill.fore_color.rgb = WHITE
    card.line.color.rgb = LIGHT
    add_rule(slide, x + 0.18, y + 0.22, 0.38, accent, 0.04)
    add_text(slide, title, x + 0.18, y + 0.42, w - 0.36, 0.3, size=14, bold=True)
    add_text(slide, body, x + 0.18, y + 0.82, w - 0.36, h - 0.95, size=10, color=GRAY)
    return card


def add_stat(slide, x: float, y: float, w: float, value: str, label: str, color=BLACK):
    add_text(slide, value, x, y, w, 0.55, size=30, bold=True, color=color, align=PP_ALIGN.CENTER)
    add_text(slide, label, x, y + 0.58, w, 0.4, size=10, color=GRAY, align=PP_ALIGN.CENTER)


def add_image_fit(slide, path: Path, x: float, y: float, w: float, h: float):
    pic = slide.shapes.add_picture(str(path), Inches(x), Inches(y), width=Inches(w))
    if pic.height > Inches(h):
        pic.height = Inches(h)
    return pic


def slide_title(prs: Presentation):
    slide = add_slide(prs, dark=True)
    add_text(slide, "BRM-Net", 0.75, 0.72, 3.2, 0.5, size=18, color=ACCENT, bold=True)
    add_text(
        slide,
        "面向多模态遥感分类的\n预算可控紧凑导出与退化监督可靠融合",
        0.75,
        1.35,
        10.8,
        1.45,
        size=30,
        bold=True,
        color=WHITE,
    )
    add_text(
        slide,
        "当前投稿版思路、证据与下一步计划 | 2026-07-20",
        0.78,
        3.25,
        8.6,
        0.35,
        size=13,
        color=LIGHT,
    )
    add_rule(slide, 0.78, 4.05, 4.1, ACCENT, 0.04)
    add_text(slide, "Fusion-Compatible Budgeted Compact Export\nwith Degradation-Supervised Reliability", 0.78, 4.28, 9.2, 0.7, size=14, color=LIGHT)
    add_text(slide, "Anonymous review draft", 9.2, 6.75, 2.6, 0.25, size=10, color=GRAY)


def slide_track(prs: Presentation):
    slide = add_slide(prs)
    add_title(slide, "当前策略：投稿优先，毕业论文扩展分轨")
    add_card(slide, 0.7, 1.5, 3.5, 3.8, "投稿主线", "紧凑导出 / 融合兼容\n可用性掩码 / 退化监督可靠性", ACCENT)
    add_card(slide, 4.55, 1.5, 3.5, 3.8, "毕业论文扩展", "动态路由、patch 级样本、\nweighted CE、演示系统、失败分析", BLUE)
    add_card(slide, 8.4, 1.5, 3.5, 3.8, "当前目标", "形成导师可审阅、可打包、\n可解释的投稿版本。", GREEN)
    add_rule(slide, 0.75, 6.25, 10.9, LIGHT, 0.02)
    add_text(slide, "核心原则：不再把所有 thesis 工作量塞入同一篇投稿稿，避免主线发散。", 0.85, 6.42, 10.4, 0.35, size=15, bold=True)


def slide_problem(prs: Presentation):
    slide = add_slide(prs)
    add_title(slide, "问题定义：轻量化与模态不稳定是耦合问题")
    items = [
        ("资源预算", "多模态分支和融合层增加 MACs、参数与推理延迟。"),
        ("缺失模态", "传感器缺失或平台不具备某一模态时，fusion softmax 需要确定性屏蔽。"),
        ("退化模态", "噪声、低分辨率、遮挡等场景下，模态仍可用但可靠性下降。"),
        ("真实导出", "软门控不能等同于部署模型，必须实际删除通道并报告资源。"),
    ]
    for i, (t, b) in enumerate(items):
        x = 0.75 + (i % 2) * 5.65
        y = 1.45 + (i // 2) * 2.05
        add_card(slide, x, y, 5.1, 1.55, t, b, [ACCENT, BLUE, GREEN, RED][i])
    add_text(slide, "因此本文的对象不是单纯剪枝，也不是重型生成式补全，而是可导出的紧凑多模态分类器。", 0.9, 6.18, 10.7, 0.48, size=15, bold=True)


def slide_method(prs: Presentation):
    slide = add_slide(prs)
    add_title(slide, "方法框架：结构门控、兼容导出、可靠融合")
    xs = [0.7, 3.35, 6.0, 8.65]
    titles = ["输入模态", "预算门控", "可靠融合", "紧凑导出"]
    bodies = [
        "HSI/main + LiDAR/aux\n可用性向量 a",
        "Hard-concrete gates\n资源目标 L_bud -> rho",
        "Quality score q_m\nMasked softmax fusion",
        "Hard thresholding\n物理删除通道",
    ]
    colors = [BLACK, ACCENT, GREEN, BLUE]
    for i, x in enumerate(xs):
        add_card(slide, x, 2.0, 2.25, 2.1, titles[i], bodies[i], colors[i])
        if i < len(xs) - 1:
            add_text(slide, "→", x + 2.38, 2.65, 0.45, 0.4, size=24, bold=True, color=ACCENT)
    add_text(slide, "关键约束：分支内通道删除不能破坏融合层所需的终端特征维度。", 1.05, 5.0, 10.2, 0.4, size=18, bold=True)
    add_text(slide, "这也是本文区别于普通单主干剪枝的主要技术叙事。", 1.05, 5.52, 10.2, 0.35, size=13, color=GRAY)


def slide_budget(prs: Presentation):
    slide = add_slide(prs)
    add_title(slide, "证据 1：目标预算与实际导出资源对齐")
    add_stat(slide, 0.85, 1.55, 2.4, "64.96%", "Target 65% MAC", ACCENT)
    add_stat(slide, 3.75, 1.55, 2.4, "80.01%", "Target 80% MAC", GREEN)
    add_stat(slide, 6.65, 1.55, 2.4, "90.07%", "Target 90% MAC", BLUE)
    add_card(slide, 9.45, 1.35, 2.55, 2.2, "解释", "三种预算均能稳定导出，与目标 MAC 基本一致。", RED)
    add_image_fit(slide, PAPER_DIR / "figures" / "generated" / "fig_brmnet_results_overview.png", 0.95, 3.35, 10.8, 2.75)
    add_text(slide, "主结论：本文报告的是 exported compact model，而不是只看 source gated model。", 0.95, 6.5, 10.5, 0.32, size=13, bold=True)


def slide_export(prs: Presentation):
    slide = add_slide(prs)
    add_title(slide, "证据 2：learned nonuniform export 不是 uniform width 替代品")
    add_card(slide, 0.75, 1.35, 3.25, 2.2, "Source vs Compact", "80% 设置下，compact full OA 为 86.90 ± 1.65，证明真实导出后仍保持性能。", ACCENT)
    add_card(slide, 4.35, 1.35, 3.25, 2.2, "Uniform-width baseline", "target-matched uniform export full OA 为 86.29 ± 0.77。", BLUE)
    add_card(slide, 7.95, 1.35, 3.25, 2.2, "差异位置", "learned export 在 main-only、downsample-4、occlusion-50 更有优势。", GREEN)
    rows = [
        ("Full OA", "86.90", "86.29", "+0.61"),
        ("Main-only", "80.04", "79.22", "+0.82"),
        ("Downsample-4", "84.15", "82.88", "+1.27"),
        ("Occlusion-50", "85.57", "85.00", "+0.57"),
    ]
    table = slide.shapes.add_table(len(rows) + 1, 4, Inches(1.2), Inches(4.1), Inches(10.1), Inches(1.7)).table
    headers = ["Metric", "Learned", "Uniform", "Delta"]
    for c, h in enumerate(headers):
        cell = table.cell(0, c)
        cell.text = h
        cell.fill.solid()
        cell.fill.fore_color.rgb = BLACK
        cell.text_frame.paragraphs[0].runs[0].font.color.rgb = WHITE
    for r, row in enumerate(rows, 1):
        for c, value in enumerate(row):
            table.cell(r, c).text = value


def slide_availability(prs: Presentation):
    slide = add_slide(prs)
    add_title(slide, "证据 3：availability mask 主要保护物理缺失模态")
    add_stat(slide, 1.1, 1.55, 2.5, "80.04", "BRM-Net HSI-only OA", GREEN)
    add_stat(slide, 4.15, 1.55, 2.5, "76.72", "w/o fusion mask", RED)
    add_stat(slide, 7.2, 1.55, 2.5, "-3.32 pp", "missing-modality drop", ACCENT)
    add_card(slide, 1.1, 3.45, 4.7, 2.35, "实验含义", "禁用 fusion mask 后：\n完整/退化可用状态变化不大；\nHSI-only 明显下降。", BLUE)
    add_card(slide, 6.25, 3.45, 4.7, 2.35, "论文表述", "mask 的核心作用：\n缺失模态不参与 softmax 竞争；\n不是所有退化状态的提升来源。", GREEN)


def slide_reliability(prs: Presentation):
    slide = add_slide(prs)
    add_title(slide, "证据 4：退化监督让可靠性分支观察 corrupted-but-available 输入")
    add_image_fit(slide, PAPER_DIR / "figures" / "generated" / "fig_controlled_corruption_reliability_curves.png", 0.75, 1.25, 11.2, 4.65)
    add_text(slide, "关键结果：adverse-state average OA 从 74.89% 提升到 84.76%，但可靠性分数只作为诊断信号，不宣称完全物理校准。", 0.8, 6.35, 10.9, 0.35, size=13, bold=True)


def slide_multidataset(prs: Presentation):
    slide = add_slide(prs)
    add_title(slide, "多数据集定位：不是 leaderboard，而是角色分工")
    rows = [
        ("Houston2013", "主实验与消融", "86.90 ± 1.65", "官方 split，证据最完整"),
        ("Trento", "外部 HSI-LiDAR 验证", "98.94 ± 0.94", "80% MAC 下保持近饱和精度"),
        ("MUUFL", "困难 stress test", "86.87 ± 1.79", "clean full OA 有下降，但 missing/degraded 更强"),
    ]
    table = slide.shapes.add_table(4, 4, Inches(0.7), Inches(1.45), Inches(11.6), Inches(2.4)).table
    headers = ["Dataset", "Role", "Full OA", "Interpretation"]
    for c, h in enumerate(headers):
        table.cell(0, c).text = h
        table.cell(0, c).fill.solid()
        table.cell(0, c).fill.fore_color.rgb = BLACK
        table.cell(0, c).text_frame.paragraphs[0].runs[0].font.color.rgb = WHITE
    for r, row in enumerate(rows, 1):
        for c, value in enumerate(row):
            table.cell(r, c).text = value
    add_card(slide, 0.9, 4.45, 5.1, 1.5, "诚实处理 MUUFL", "不把 MUUFL 写成 clean OA 全面胜利，而是写成鲁棒性 trade-off 证据。", RED)
    add_card(slide, 6.4, 4.45, 5.1, 1.5, "导师需决策", "MUUFL 保留主文可体现工作量，但会增加解释成本。", ACCENT)


def slide_readiness(prs: Presentation):
    slide = add_slide(prs)
    add_title(slide, "投稿准备状态：结构、引用、源码包均已过检查")
    add_stat(slide, 0.85, 1.5, 2.2, "PASS", "Readiness scanner", GREEN)
    add_stat(slide, 3.35, 1.5, 2.2, "0", "Package findings", GREEN)
    add_stat(slide, 5.85, 1.5, 2.2, "20", "Clean source files", ACCENT)
    add_stat(slide, 8.35, 1.5, 2.2, "10", "PDF pages", BLUE)
    add_card(slide, 0.9, 3.55, 3.3, 1.8, "已完成", "related-work audit、readiness gate、source package、PDF visual review。", GREEN)
    add_card(slide, 4.6, 3.55, 3.3, 1.8, "当前产物", "paper.pdf 与 brmnet_pricai2026_submission_source_20260719.zip。", ACCENT)
    add_card(slide, 8.3, 3.55, 3.3, 1.8, "剩余工作", "最终语言润色、导师决策、必要时补强 baseline。", BLUE)


def slide_risks(prs: Presentation):
    slide = add_slide(prs)
    add_title(slide, "当前风险与处理方式")
    risks = [
        ("创新性风险", "不包装成大架构，强调真实 compact export + 多状态验证闭环。"),
        ("可靠性解释", "避免说完全校准，只说 degradation-supervised diagnostic signal。"),
        ("MUUFL 负面结果", "保留 trade-off 表述，不隐藏 clean full OA 下降。"),
        ("范围漂移", "routing、weighted CE、demo system 留给毕业论文扩展。"),
    ]
    for i, (t, b) in enumerate(risks):
        add_card(slide, 0.8 + (i % 2) * 5.6, 1.35 + (i // 2) * 2.15, 5.0, 1.55, t, b, [RED, ACCENT, BLUE, GREEN][i])


def slide_advisor(prs: Presentation):
    slide = add_slide(prs)
    add_title(slide, "需要导师拍板的 4 个问题")
    qs = [
        "是否接受当前题目，还是改短为 Fusion-Compatible Compact Export for Reliable Multimodal Remote Sensing Classification？",
        "MUUFL 是否保留在主文，还是移到补充/毕业论文以降低解释成本？",
        "当前是否停止加小实验，进入语言润色和投稿包整理？",
        "若仍需补强创新性，优先补 larger-backbone transfer 还是 stronger lightweight baseline？",
    ]
    for i, q in enumerate(qs, 1):
        y = 1.25 + (i - 1) * 1.25
        add_text(slide, f"{i}", 0.85, y, 0.35, 0.35, size=18, bold=True, color=ACCENT)
        add_text(slide, q, 1.35, y, 10.5, 0.55, size=15, bold=True if i == 3 else False)


def slide_next(prs: Presentation):
    slide = add_slide(prs, dark=True)
    add_text(slide, "下一步执行计划", 0.75, 0.65, 6.8, 0.55, size=30, bold=True, color=WHITE)
    add_rule(slide, 0.78, 1.42, 3.4, ACCENT, 0.04)
    steps = [
        ("1", "导师审阅", "确认题目、MUUFL 展示方式、是否补实验。"),
        ("2", "语言润色", "压缩 Method/Experiments 重复解释，统一贡献表述。"),
        ("3", "投稿包冻结", "PDF + clean source zip + README + checklist。"),
        ("4", "毕业论文扩展", "回到 routing、demo system、class-wise diagnostics。"),
    ]
    for i, (num, title, body) in enumerate(steps):
        x = 0.85 + (i % 2) * 5.55
        y = 2.0 + (i // 2) * 1.75
        add_text(slide, num, x, y, 0.45, 0.45, size=22, bold=True, color=ACCENT)
        add_text(slide, title, x + 0.6, y, 4.6, 0.35, size=18, bold=True, color=WHITE)
        add_text(slide, body, x + 0.6, y + 0.45, 4.5, 0.5, size=12, color=LIGHT)
    add_text(slide, "当前状态：已经具备导师审阅版，不建议继续无边界加模块。", 0.9, 6.45, 10.8, 0.35, size=15, bold=True, color=WHITE)


def build_deck(template: Path, output: Path) -> None:
    prs = Presentation(str(template))
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    clear_template_slides(prs)
    slide_title(prs)
    slide_track(prs)
    slide_problem(prs)
    slide_method(prs)
    slide_budget(prs)
    slide_export(prs)
    slide_availability(prs)
    slide_reliability(prs)
    slide_multidataset(prs)
    slide_readiness(prs)
    slide_risks(prs)
    slide_advisor(prs)
    slide_next(prs)
    output.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(output))


def main(argv: list[str] | None = None) -> int:
    args = build_parser().parse_args(argv)
    build_deck(Path(args.template), Path(args.output))
    print(f"Wrote deck to {args.output}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
